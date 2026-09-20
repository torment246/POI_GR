#!/usr/bin/env python3
"""Revise the uploaded report without redesigning its master or existing layouts."""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
import json
import os
from lxml import etree

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from build_0907 import Deck, style_frame
from update_0917 import (
    BASELINES, FIGURES, METHODS, ROOT, SID_ROOT, SOURCE, digest, read_json,
)

WORK = ROOT / 'qg_prqk/outputs/ppt_0917/original_style'
OUTPUT = ROOT / 'PPT/胡丹-0917-QG-HRQ原版式修订.pptx'
NAMES = ['PRQ-KMeans', 'QG-HRQ']


def replace_text(shape, text: str) -> None:
    """Keep the existing geometry, paragraph formatting and first-run appearance."""
    frame = shape.text_frame
    template = deepcopy(frame.paragraphs[0]._p)
    for p in list(frame._txBody.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}p')):
        frame._txBody.remove(p)
    for line in text.split('\n'):
        element = deepcopy(template)
        frame._txBody.append(element)
        paragraph = frame.paragraphs[-1]
        runs = list(paragraph.runs)
        if runs:
            runs[0].text = line
            for run in runs[1:]:
                element.remove(run._r)
        else:
            paragraph.add_run().text = line


def matching(slide, prefix: str):
    return next(s for s in slide.shapes if s.has_text_frame and s.text.startswith(prefix))


class Revision(Deck):
    def __init__(self) -> None:
        self.prs = Presentation(SOURCE)
        self.original = list(self.prs.slides)
        self.layout = self.original[11].slide_layout
        self.title = deepcopy(self.original[11].shapes[0]._element)
        reference = Presentation(ROOT / 'PPT/胡丹-0810.pptx')
        self.table_props = deepcopy(next(s.table._tbl.tblPr for s in reference.slides[7].shapes if s.has_table))
        self.sources = {}

    def page(self, title: str):
        slide = self.prs.slides.add_slide(self.layout)
        for shape in list(slide.shapes):
            slide.shapes._spTree.remove(shape._element)
        slide.shapes._spTree.insert_element_before(deepcopy(self.title), 'p:extLst')
        replace_text(slide.shapes[0], title)
        slide.shapes[0].width = Inches(12.40)
        return slide

    def results_page(self, slide, title: str) -> None:
        for shape in list(slide.shapes):
            if shape.shape_type == 13:
                slide.shapes._spTree.remove(shape._element)
        heading = next(s for s in slide.shapes if s.has_text_frame)
        replace_text(heading, title)
        heading.width = Inches(12.40)

    def table(self, slide, headers, rows, widths, **kwargs):
        end = super().table(slide, headers, rows, widths, highlight=None, **kwargs)
        table = next(s.table for s in slide.shapes if s.has_table)
        for row in list(table.rows)[1:]:
            for cell in row.cells:
                cell.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = False
        return end

    def evidence(self, path: Path) -> dict:
        self.sources[str(path.relative_to(ROOT))] = digest(path)
        return read_json(path)


def revise_text(deck: Revision) -> None:
    """Change wording in place; the original method flows remain editable."""
    replacements = {
        1: {
            '问题一：': '问题一：POI 内容相似\n不等于 Query 指向相同目标',
            '问题二：': '问题二：同类、同名 POI\n仍需要结合位置做实体区分',
            'QG-HRQ': 'Query 粒度监督',
            '在离线 SID': '粗类 Query 对齐粗前缀，\n精确 Query 才约束完整 SID。',
            'BeamRisk-SFT': '类别与局部地理',
            '固定 SID': '前两层组织类别语义，\n末层在地理父组内区分实体。',
        },
        2: {
            'SID 目标必须': 'SID 构建需考虑 Query 可预测性',
            '保留独立 Query': '独立 Query 视图，共享离散索引',
            'Geo 只在': 'Geo 在 GID/S1/S2 父组内约束 S3',
        },
        3: {
            '硬 category index': '首层硬类别；未按 Query 歧义\n分配逐层监督深度',
            '绝对地理解决': '已建模地理，但未按 Query\n目标分布定义监督深度',
            '相关论文都有': 'QG-HRQ：按 Query 粒度分层对齐，将类别语义与局部地理分工建模',
        },
        4: {
            '残留碰撞追加': '仅残留碰撞追加 Dedup\nFinal ID = GID6 + SID3 + [D]',
            'D1→': 'D1→S1；D2→S1/S2；D3→S1/S2/S3。\n类别约束 S1/S2；Geo 在 GID6/S1/S2 父组内约束 S3。',
        },
        5: {'同一个 Query': '按 Train 目标分布的集中度与熵确定监督深度；下方 Query 为机制示意。'},
        8: {'一个粗类': '同一粗类可占多个 S1 码字，不硬绑定单一节点。'},
        10: {
            '每层从': '先构建 POI-only PRQ-KMeans，再用其对应层初始化，交替更新双视图分配与中心。',
            '目标相对改善': '目标相对改善 <1e-4；POI 变化 <0.1%；Query <0.2%；\n连续两轮满足，或最多 60 轮。',
        },
        18: {'生成式 POI': 'QG-HRQ 在 POI-only 量化基础上，引入 Query 粒度监督、类别软锚点与局部地理约束。\n\n前缀类别纯度和碰撞指标已有改善；继续通过同协议 SFT 评测与消融验证检索收益。'},
    }
    for index, changes in replacements.items():
        for old, new in changes.items():
            replace_text(matching(deck.original[index], old), new)
    matching(deck.original[18], 'QG-HRQ 在').height = Inches(3.8)
    for label, url in [
        ('TIGER', 'https://arxiv.org/abs/2305.05065'),
        ('PRQ-KMeans', 'https://arxiv.org/abs/2608.24207'),
        ('CQ-SID', 'https://arxiv.org/abs/2605.14434'),
        ('GenPOI', 'https://arxiv.org/abs/2605.03397'),
    ]:
        matching(deck.original[3], label).text_frame.paragraphs[0].runs[0].hyperlink.address = url

    # The upload has arrows but lacks the corresponding module labels.
    template = matching(deck.original[8], '当前 residual')
    def node(slide, label, x, y, w, h):
        slide.shapes._spTree.insert_element_before(deepcopy(template._element), 'p:extLst')
        shape = slide.shapes[-1]
        shape._element.nvSpPr.cNvPr.set('id', str(slide.shapes._next_shape_id))
        shape.left, shape.top, shape.width, shape.height = map(Inches, (x, y, w, h))
        replace_text(shape, label)
        return shape
    for label, x, y, w, h in [
        ('POI residual\n内容视图', .65, 2.10, 2.70, 1.02),
        ('Query residual\n查询视图', .65, 4.40, 2.70, 1.02),
        ('POI 质心 U\nPOI 分配', 4.02, 2.10, 2.60, 1.02),
        ('Query 质心 V\nQuery 分配', 4.02, 4.40, 2.60, 1.02),
        ('共享离散 token k\n图约束促进一致', 7.40, 3.20, 2.32, 1.16),
    ]:
        node(deck.original[7], label, x, y, w, h)
    node(deck.original[8], '选中码向量 c_k', 10.30, 2.68, 2.27, .72)
    s = deck.original[10]
    node(s, 'POI-only\n码本初始化', .65, 2.04, 1.64, 1.12)
    node(s, '分层 Query\n与 POI 图', 2.73, 2.04, 1.64, 1.12)
    s.shapes[11].left = Inches(9.76)
    s.shapes[12].width = Inches(9.76 - 4.78)


def visualizations(deck: Revision) -> list:
    """Replot frozen numerical data; do not change embeddings or fit t-SNE."""
    import numpy as np
    import pyarrow.parquet as pq
    os.environ.setdefault('MPLCONFIGDIR', str(WORK / 'mpl_cache'))
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.patches import Patch

    font_path = ROOT / 'qg_prqk/outputs/inputs/fonts/NotoSansCJKsc-Regular.otf'
    font_manager.fontManager.addfont(str(font_path))
    plt.rcParams.update({'font.family': font_manager.FontProperties(fname=str(font_path)).get_name(),
                         'axes.unicode_minus': False, 'font.size': 18})
    path = FIGURES / 'qg_prqk_sid_category_region_v2/sample_and_coordinates.parquet'
    deck.sources[str(path.relative_to(ROOT))] = digest(path)
    sample = pq.read_table(path).to_pydict()
    colors = ['#3977b8', '#e79d30', '#309e88', '#c66594', '#8264b4']
    categories = ['房产小区', '室内及附属设施', '美食', '购物', '生活服务']
    coordinates = np.asarray([np.column_stack([sample[f'{m}_tsne_x'], sample[f'{m}_tsne_y']]) for m in METHODS])
    low, high = coordinates.min(axis=(0, 1)), coordinates.max(axis=(0, 1))
    margin = (high - low) * .05
    slides = [deck.page('SID 可视化：同一批五类 POI')]
    for i, label in enumerate(NAMES):
        figure, ax = plt.subplots(figsize=(6, 5.3))
        for category, color in zip(categories, colors):
            mask = np.asarray(sample['coarse_category']) == category
            ax.scatter(*coordinates[i, mask].T, s=19, c=color, alpha=.78, linewidths=0)
        ax.set_xlim(low[0] - margin[0], high[0] + margin[0])
        ax.set_ylim(low[1] - margin[1], high[1] + margin[1])
        ax.set_aspect('equal'); ax.axis('off')
        figure.subplots_adjust(left=0, right=1, bottom=0, top=1)
        path = WORK / f'tsne_{i}.png'
        figure.savefig(path, dpi=180, transparent=True); plt.close(figure)
        deck.text(slides[0], label + ('（POI-only）' if i == 0 else '（带 GID）'), 1.20, .45,
                  x=.65 + i * 6.25, w=5.8, size=21, align=PP_ALIGN.CENTER)
        slides[0].shapes.add_picture(str(path), Inches(.65 + i * 6.25), Inches(1.78), width=Inches(5.80))
    # Legends are ordinary editable slide labels, not embedded explanatory text.
    from pptx.enum.shapes import MSO_SHAPE
    for i, (category, color) in enumerate(zip(categories, colors)):
        x = .85 + i * 2.5
        dot = slides[0].shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(6.65), Inches(.12), Inches(.12))
        dot.fill.solid(); dot.fill.fore_color.rgb = RGBColor.from_string(color[1:]); dot.line.fill.background()
        deck.text(slides[0], category, 6.55, .4, x=x + .18, w=2.25, size=17)

    cat = deck.evidence(FIGURES / 'qg_prqk_sid_category_region_v2/metrics.json')
    slides.append(s := deck.page('SID 前缀类别纯度（全库，未拼接 GID / Dedup）'))
    figure, ax = plt.subplots(figsize=(12.2, 5.8))
    x = np.arange(3)
    for i, (method, color) in enumerate(zip(METHODS[:2], ['#9BAFC4', '#3977b8'])):
        values = [100 * v['fine_category']['poi_weighted_top1_share'] for v in cat['methods'][method]['catalog_prefix_metrics']]
        bars = ax.bar(x + (i - .5) * .30, values, .30, label=NAMES[i], color=color)
        ax.bar_label(bars, labels=[f'{v:.2f}' for v in values], padding=6, fontsize=18)
    ax.set_xticks(x, ['S1', 'S1 / S2', 'S1 / S2 / S3']); ax.set_ylim(0, 110)
    ax.set_yticks([0, 25, 50, 75, 100]); ax.set_ylabel('Micro Purity (%)')
    ax.legend(frameon=False, loc='upper left', ncol=2)
    ax.spines[['top', 'right']].set_visible(False)
    figure.tight_layout()
    path = WORK / 'prefix_purity.png'; figure.savefig(path, dpi=180); plt.close(figure)
    s.shapes.add_picture(str(path), Inches(.62), Inches(1.25), width=Inches(12.06))

    collision = deck.evidence(FIGURES / 'qg_prqk_sid_collision_examples_v1/collision_examples.json')
    slides.append(s := deck.page('SID 前缀分布：自行车专卖碰撞案例'))
    figure, axes = plt.subplots(2, 3, figsize=(16.8, 8.2))
    for depth in range(1, 4):
        examples = [next(e for e in collision['methods'][m]['prefix_examples'] if e['anchor_index'] == 3 and e['level'] == depth) for m in METHODS[:2]]
        for row, key in enumerate(['fine_category', 'region_geohash5']):
            ax = axes[row, depth - 1]
            counts = [e[key]['counts'] for e in examples]
            # Shared named labels, not per-method ranks; aggregate only the tail.
            labels = sorted(set().union(*(c.keys() for c in counts)), key=lambda k: (-sum(c.get(k, 0) for c in counts), k))[:3]
            palette = ['#3977b8', '#80add6', '#b8d4eb', '#E2E8F0'] if row == 0 else ['#d78a2d', '#e6b46f', '#f2d9b3', '#E2E8F0']
            for i, e in enumerate(examples):
                vals = [100 * counts[i].get(k, 0) / e['poi_count'] for k in labels]
                vals.append(100 - sum(vals)); left = 0
                for value, color in zip(vals, palette):
                    ax.barh(1 - i, value, left=left, height=.40, color=color); left += value
            ax.set_yticks([1, 0], [NAMES[i] + '\n[' + ','.join(map(str, e['prefix'])) + ']' + f"  n={e['poi_count']}" for i, e in enumerate(examples)], fontsize=13)
            ax.set_xlim(0, 100); ax.set_xticks([0, 50, 100], ['0%', '50%', '100%'], fontsize=14)
            ax.set_ylim(-.5, 1.6)
            ax.set_title(['S1', 'S1 / S2', 'S1 / S2 / S3'][depth - 1] + (' · 类别' if row == 0 else ' · Geohash5'), fontsize=19, pad=14)
            ax.spines[['top', 'right', 'left']].set_visible(False)
            ax.tick_params(axis='y', length=0)
            legend = [k.split(':')[-1] for k in labels] + ['其他']
            ax.legend(handles=[Patch(color=c, label=l) for c, l in zip(palette, legend)], loc='upper center', bbox_to_anchor=(.40, -.16), ncol=2, frameon=False, fontsize=13)
    figure.subplots_adjust(left=.115, right=.965, top=.91, bottom=.15, wspace=.95, hspace=.91)
    path = WORK / 'prefix_composition.png'; figure.savefig(path, dpi=180); plt.close(figure)
    s.shapes.add_picture(str(path), Inches(.30), Inches(1.12), width=Inches(12.70))
    return slides


def result_tables(deck: Revision) -> None:
    s = deck.original[11]
    deck.results_page(s, 'SID 构建指标（512×3，未拼接 GID / Dedup）')
    rows = [
        ['TIGER', 'RQ-VAE', '77.03%', '33.78%', '152', '59.29 / 70.17 / 95.82'],
        ['GNPR', 'RQ-VAE + diversity', '85.77%', '22.56%', '261', '64.06 / 76.04 / 98.82'],
        ['GenPOI', 'RQ-VAE', '75.18%', '36.58%', '191', '55.00 / 68.63 / 93.60'],
        ['MMBERT', 'RQ-VAE', '90.60%', '16.22%', '30', '45.72 / 63.71 / 97.17'],
    ]
    stat = deck.evidence(SID_ROOT / 'p8_nogid_s1s2_parent_comparison_v1/full_342879q_716245p/static_metrics.json')
    cat = deck.evidence(FIGURES / 'qg_prqk_sid_category_region_v2/metrics.json')
    for m, name in zip(METHODS, ['PRQ-KMeans\n(POI-only)', 'QG-HRQ', 'QG-HRQ\nw/o GID（消融）']):
        sid = stat['methods'][m]['paths']['s1_s2_s3']
        levels = cat['methods'][m]['catalog_prefix_metrics']
        rows.append([name, 'PRQ-KMeans', f"{sid['distinct_ratio_of_pois']*100:.2f}%", f"{(1-levels[2]['singleton_poi_share'])*100:.2f}%", str(int(sid['bucket']['max'])), ' / '.join(f"{l['fine_category']['poi_weighted_top1_share']*100:.2f}" for l in levels)])
    deck.table(s, ['方法', 'Quantizer', '唯一 SID\n比例', '碰撞 POI\n比例', '最大桶', 'D1 / D2 / D3\nMicro Purity (%)'], rows,
               [2.35, 2.0, 1.3, 1.3, .85, 4.26], y=1.24, size=15, min_row=.52)
    titles = ['固定 10k Validation', '已见 Query / 未见 Query–POI', '新 Query / 已见目标', '长尾目标', '冷目标', '全量 Test（606,682 条）']
    for index, (key, title) in enumerate(zip(BASELINES, titles)):
        test = key == 'test'; s = deck.original[12 + index]
        deck.results_page(s, '实验结果：' + title)
        rows = [[name, '无约束', *[f'{v:.2f}' for v in values]] for name, values in zip(['TIGER', 'MMBERT'], BASELINES[key])]
        for constrained in ([False] if test else [False, True]):
            run = 'sft_epoch3_full_test_20260714_v1' if test else 'sft_epoch3_constrained_fixed10k_generalization_v1' if constrained else 'sft_epoch3_fixed10k_generalization_v1'
            for variant, name in [('a4_gid_parent', 'QG-HRQ'), ('a4_nogid', 'QG-HRQ\nw/o GID（消融）')]:
                path = ROOT / 'qg_prqk/outputs/eval' / run / variant / 'results'
                result = deck.evidence(path / 'result.json' if test else path / key / 'result.json')
                assert result['status'] == 'completed'
                metrics = result['metrics']; assert metrics['sample_count'] == (606682 if test else 10000)
                rows.append([name, '约束' if constrained else '无约束', *[f'{metrics[m]*100:.2f}' for m in ['hr@1', 'hr@3', 'hr@5', 'hr@10', 'ndcg@10', 'valid_id_rate']]])
        deck.table(s, ['方法', '解码', 'HR@1', 'HR@3', 'HR@5', 'HR@10', 'NDCG@10', 'Valid ID'], rows,
                   [2.65, 1.05, 1.22, 1.22, 1.22, 1.22, 1.75, 1.73], y=1.60, size=16, min_row=.63)


def validate_output(output: Path, source_hash: str) -> None:
    """Compare XML semantics, ignoring serialization-only XML declaration changes."""
    with ZipFile(SOURCE) as before, ZipFile(output) as after:
        assert after.testzip() is None
        protected = [n for n in before.namelist() if n.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/'))
                     or (n.startswith('ppt/notesSlides/') and '/_rels/' not in n)]
        for name in protected:
            original_root = etree.fromstring(before.read(name))
            revised_root = etree.fromstring(after.read(name))
            if name.endswith('.rels'):
                # Relationship ordering has no semantic meaning in OPC packages.
                original_root[:] = sorted(original_root, key=lambda e: e.get('Id'))
                revised_root[:] = sorted(revised_root, key=lambda e: e.get('Id'))
            original_xml = etree.tostring(original_root, method='c14n')
            revised_xml = etree.tostring(revised_root, method='c14n')
            assert original_xml == revised_xml, f'模板或原备注被修改：{name}'
    result = Presentation(output)
    assert len(result.slides) == 22
    assert digest(SOURCE) == source_hash
    text = '\n'.join(s.text for p in result.slides for s in p.shapes if s.has_text_frame)
    assert not any(token in text for token in ['A0', 'A4', 'BeamRisk', '来源：'])
    for page in result.slides:
        for shape in page.shapes:
            if not shape.has_table: continue
            for row in list(shape.table.rows)[1:]:
                assert all(str(c.fill.fore_color.rgb) == 'FFFFFF' for c in row.cells)
                assert all(not r.font.bold for c in row.cells for p in c.text_frame.paragraphs for r in p.runs)
    assert sum(s.has_table for p in result.slides for s in p.shapes) == 7
    print('已核验：22 页、原母版/布局/主题/备注不变、7 页原生表格、数据行统一底色与字重。')


def build(output: Path) -> None:
    if output.exists() or output.resolve() == SOURCE.resolve():
        raise ValueError('请使用不存在的新输出文件；原稿和旧版本均不覆盖。')
    WORK.mkdir(parents=True, exist_ok=True)
    source_hash = digest(SOURCE)
    deck = Revision()
    assert len(deck.original) == 19
    revise_text(deck); added = visualizations(deck); result_tables(deck)
    order = [*deck.original[:11], *added, *deck.original[11:]]
    ids = deck.prs.slides._sldIdLst
    lookup = {item.id: item for item in ids}
    ordered_ids = [lookup[s.slide_id] for s in order]
    for item in list(ids): ids.remove(item)
    for item in ordered_ids: ids.append(item)
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.prs.save(output)
    validate_output(output, source_hash)
    manifest = {'source_sha256': source_hash, 'output': str(output), 'output_sha256': digest(output),
                'slides': 22, 'masters_layouts_themes_notes_unchanged': True, 'native_result_tables': 7,
                'metric_sources': deck.sources, 'tsne_refit': False,
                'tsne_scope': 'Two panels selected from the unchanged frozen joint three-method projection.',
                'collision_case': 'Frozen seed-42 shopping anchor; shared top-3 labels plus other; no external GID in grouping.',
                'style_override': 'User requests original white titles on existing blue header; no template recoloring.'}
    (WORK / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'output': str(output), 'slides': 22, 'original_style_preserved': True}, ensure_ascii=False))


def replace_s3_visualizations(source: Path, figures: Path, output: Path) -> None:
    """Replace only SID evidence pages, keeping every other slide unchanged."""
    if output.exists() or output.resolve() == source.resolve():
        raise ValueError('请使用不存在的新输出文件；原稿和旧版本均不覆盖。')
    manifest = read_json(figures / 'manifest.json')
    if manifest.get('status') != 'completed':
        raise ValueError('S3 可视化尚未完成')
    for name, expected in manifest['artifacts'].items():
        if digest(figures / name) != expected:
            raise ValueError(f'S3 图表产物哈希不匹配：{name}')
    data = read_json(figures / 'analysis.json')
    if data['selection']['selection_uses_final_s3'] or len(data['cases']) != 2:
        raise ValueError('案例必须是预先确定的两个父组，不能按最终改进选例')
    source_hash = digest(source)
    prs = Presentation(source)
    original = list(prs.slides)
    if len(original) != 22 or not original[11].shapes[0].text.startswith('SID 可视化'):
        raise ValueError('输入应为已确认的 22 页原版式修订稿')

    def body(slide, title):
        for shape in list(slide.shapes)[1:]:
            slide.shapes._spTree.remove(shape._element)
        replace_text(slide.shapes[0], title)

    def label(slide, text, y, h=.42, *, x=.65, w=12.05, size=20, align=PP_ALIGN.LEFT):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        style_frame(shape.text_frame, text, size, '111827', False, align)

    overview = original[11]
    body(overview, 'S3 更新前后：固定 GID / S1 / S2')
    label(overview, '全库 716,245 个 POI · 同一父组划分 · 仅比较 S3 分配', 1.16, size=21)
    overview.shapes.add_picture(str(figures / 's3_local_overview.png'), Inches(.64), Inches(1.86), width=Inches(12.05))

    extra = prs.slides.add_slide(original[13].slide_layout)
    for shape in list(extra.shapes):
        extra.shapes._spTree.remove(shape._element)
    extra.shapes._spTree.insert_element_before(deepcopy(original[13].shapes[0]._element), 'p:extLst')
    case_slides = [original[13], extra]
    # These are display labels for the frozen seed-42 selection, not search criteria.
    case_titles = {62206: '星光视界中心', 384796: '奶东村楼栋'}
    for index, (slide, case) in enumerate(zip(case_slides, data['cases']), 1):
        title = case_titles.get(case['parent_group'], f"父组 {case['parent_group']}")
        body(slide, 'S3 局部编码：' + title)
        label(slide, f"固定父组：{case['gid6']} / {case['s1']} / {case['s2']}", 1.10, w=6.0, size=20)
        label(slide, '同色 / 连线 = 同 S3；编号 = 同一 POI', 1.10, x=6.72, w=6.0, size=20)
        slide.shapes.add_picture(str(figures / f's3_local_case_{index}.png'), Inches(.65), Inches(1.56), width=Inches(12.05))
        for column, key in enumerate(('before_collision_pairs', 'after_collision_pairs')):
            label(slide, f"同码 POI 对数：{case[key]}", 5.88, x=.85 + column * 6.2, w=5.65,
                  size=21, align=PP_ALIGN.CENTER)
        # Number-to-name mapping is a readable figure legend, not a footnote.
        members = case['members']
        for i, member in enumerate(members):
            column, row = divmod(i, 3)
            label(slide, f"{i + 1}  {member['displayname']}", 6.35 + row * .33, h=.33,
                  x=.70 + column * 6.18, w=6.08, size=16)

    order = [*original[:11], original[12], overview, *case_slides, *original[14:]]
    ids = prs.slides._sldIdLst
    lookup = {item.id: item for item in ids}
    ordered_ids = [lookup[slide.slide_id] for slide in order]
    for item in list(ids):
        ids.remove(item)
    for item in ordered_ids:
        ids.append(item)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)

    result = Presentation(output)
    assert len(result.slides) == 23
    by_id = {s.slide_id: s for s in result.slides}
    baseline = Presentation(source)
    for index, slide in enumerate(baseline.slides):
        revised = by_id[slide.slide_id]
        if index not in (11, 13):
            assert etree.tostring(slide._element, method='c14n') == etree.tostring(revised._element, method='c14n'), f'非可视化页面被改动：{index + 1}'
        assert slide.has_notes_slide == revised.has_notes_slide
        if slide.has_notes_slide:
            assert etree.tostring(slide.notes_slide._element, method='c14n') == etree.tostring(revised.notes_slide._element, method='c14n')
    with ZipFile(source) as before, ZipFile(output) as after:
        assert after.testzip() is None
        for name in before.namelist():
            if not name.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/')):
                continue
            roots = [etree.fromstring(z.read(name)) for z in (before, after)]
            if name.endswith('.rels'):
                for root in roots:
                    root[:] = sorted(root, key=lambda e: e.get('Id'))
            assert etree.tostring(roots[0], method='c14n') == etree.tostring(roots[1], method='c14n'), f'母版/布局/主题被修改：{name}'
    assert digest(source) == source_hash
    tables = [s.table for p in result.slides for s in p.shapes if s.has_table]
    assert len(tables) == 7
    for table in tables:
        for row in list(table.rows)[1:]:
            assert all(str(c.fill.fore_color.rgb) == 'FFFFFF' for c in row.cells)
            assert all(not r.font.bold for c in row.cells for p in c.text_frame.paragraphs for r in p.runs)
    record = {'source': str(source), 'source_sha256': source_hash, 'output': str(output), 'output_sha256': digest(output),
              'figure_manifest_sha256': digest(figures / 'manifest.json'), 'comparison': data['comparison'],
              'selection': data['selection'], 'limits': data['limits'], 'slides': 23,
              'unchanged_source_slides': 20, 'native_result_tables': 7,
              'masters_layouts_themes_notes_unchanged': True,
              'style_override': '用户要求保留原稿蓝色页眉与白色标题；不重新设计版式。'}
    record_dir = figures / output.stem
    record_dir.mkdir(exist_ok=False)
    (record_dir / 'ppt_manifest.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'output': str(output), 'slides': 23, 'unchanged_source_slides': 20, 'native_result_tables': 7}, ensure_ascii=False))


def append_query_visualizations(source: Path, analysis_dir: Path, output: Path) -> None:
    """Append query trees and an editable statistics table to the accepted deck."""
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Pt

    if output.exists() or output.resolve() == source.resolve():
        raise ValueError('请指定新的输出 PPT，不覆盖任何旧稿。')
    manifest = read_json(analysis_dir / 'manifest.json')
    if manifest.get('status') != 'completed':
        raise ValueError('Query 分析未完成')
    for name, expected in manifest['artifacts'].items():
        if digest(analysis_dir / name) != expected:
            raise ValueError(f'Query 分析哈希不匹配：{name}')
    data = read_json(analysis_dir / 'analysis.json')
    source_hash = digest(source)
    deck = Revision.__new__(Revision)
    deck.prs = Presentation(source)
    original = list(deck.prs.slides)
    if len(original) != 23:
        raise ValueError('Query 增页模式需要已确认的 23 页 S3 局部消歧版')
    deck.layout = original[11].slide_layout
    deck.title = deepcopy(original[11].shapes[0]._element)
    deck.table_props = deepcopy(next(s.table._tbl.tblPr for p in original for s in p.shapes if s.has_table))
    positive = data['selection'].get('case_selection') == 'positive'
    contrast = data['selection'].get('case_selection') in ('contrast', 'positive')
    unchanged_cases = all(abs(c['before']['top_prefix_share'] - c['after']['top_prefix_share']) < 1e-12 for c in data['cases'])
    title = ('Query 引导：相关目标共享前缀的改善样例' if positive else
             'Query 前缀变化：改善与退化案例' if contrast else
             'Query 前缀案例：随机样例的集中度未变' if unchanged_cases else
             'Query 的目标 POI 分布在哪些前缀？')
    tree = deck.page(title)
    if contrast:
        selection_text = ('从 D1 / D2 改善样例中以 seed=42 抽样；用于解释机制，不代表平均效果。' if positive else
                          '按改善 / 退化分组，组内 seed=42 抽样；用于解释变化，不代表平均效果。')
        deck.text(tree, selection_text,
                  1.02, .40, size=19)
    for column, name in enumerate(('PRQ-KMeans（POI-only）', 'QG-HRQ')):
        deck.text(tree, name, 1.49 if contrast else 1.04, .42, x=.67 + column * 6.18, w=5.80, size=21, align=PP_ALIGN.CENTER)

    def line(slide, x1, y1, x2, y2, weight=1.0):
        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shape.line.color.rgb = RGBColor.from_string('64748B')
        shape.line.width = Pt(weight)

    def dot(slide, x, y, diameter):
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - diameter / 2), Inches(y - diameter / 2), Inches(diameter), Inches(diameter))
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string('1A3A6B')
        shape.line.fill.background()

    for row, case in enumerate(data['cases']):
        top = 2.00 + row * 2.55 if contrast else 1.58 + row * 2.75
        spacing = .53 if contrast else .55
        level = 'S1' if case['depth'] == 1 else 'S1/S2'
        if contrast:
            outcome = '改善' if case['selection_outcome'] == 'improved' else '退化'
            deck.text(tree, f"{outcome}案例  D{case['depth']}：{case['query']}（{level}）", top, .38, x=.69, w=7.0, size=18)
            deck.text(tree, f"主前缀占比 {case['before']['top_prefix_share']:.1%} → {case['after']['top_prefix_share']:.1%}",
                      top, .38, x=8.12, w=4.55, size=18, align=PP_ALIGN.RIGHT)
        else:
            subtitle = f"D{case['depth']}  Query：{case['query']}   |   {level} 前缀   |   {case['retained_orders']} 条有效订单"
            deck.text(tree, subtitle, top, .38, x=.69, w=11.94, size=18)
        for column, variant in enumerate(('before', 'after')):
            x = .70 + column * 6.18
            groups = {}
            for member in case['members']:
                groups.setdefault(tuple(member[f'{variant}_prefix']), []).append(member)
            groups = sorted(groups.items(), key=lambda item: (-sum(m['pair_count'] for m in item[1]), item[0]))
            cursor, positions = 0, []
            for code, members in groups:
                members.sort(key=lambda m: m['poi_row_index'])
                ys = [top + .56 + (cursor + j) * spacing for j in range(len(members))]
                positions.append((code, members, ys, sum(ys) / len(ys)))
                cursor += len(members)
            root_y = top + .56 + (cursor - 1) * spacing / 2
            for code, members, ys, center in positions:
                share = sum(m['pair_count'] for m in members) / case['retained_orders']
                line(tree, x + .17, root_y, x + 1.05, center, 1 + share * 2)
                for member, y in zip(members, ys):
                    line(tree, x + 2.18, center, x + 2.65, y, 1)
            dot(tree, x + .17, root_y, .12)
            for code, members, ys, center in positions:
                share = sum(m['pair_count'] for m in members) / case['retained_orders']
                prefix = '/'.join(map(str, code))
                deck.text(tree, f'{prefix}\n{share:.1%}', center - .25, .55,
                          x=x + 1.02, w=1.24, size=16, align=PP_ALIGN.CENTER)
                for member, y in zip(members, ys):
                    # The complete name and count identify each unchanged target.
                    deck.text(tree, f"{member['displayname']}  ×{member['pair_count']}", y - .23, .55,
                              x=x + 2.70, w=3.00, size=16)

    stats = deck.page('Query 目标前缀集中度：全量 Train 对比')
    deck.text(stats, '有效多目标 Query；Query 等权；不拼接 GID / Dedup', 1.17, .45, size=21)
    rows = []
    for depth in (1, 2):
        metrics = data['metrics'][f'D{depth}']
        for variant, name in (('before', 'PRQ-KMeans'), ('after', 'QG-HRQ')):
            values = metrics[variant]
            rows.append([f"D{depth} / {'S1' if depth == 1 else 'S1/S2'}\n{name}",
                         f"{metrics['multi_target_queries']:,}", f"{values['top_prefix_share'] * 100:.2f}%",
                         f"{values['prefix_count']:.2f}", f"{values['entropy_bits']:.3f}",
                         f"{values['expected_bucket_size']:,.2f}"])
    end = deck.table(stats, ['分组 / 方法', 'Query 数', '主前缀\n目标占比', '平均前缀数', '前缀熵\n(bits)', '加权桶大小\n(POI 数)'],
                     rows, [3.15, 1.35, 1.95, 1.70, 1.55, 2.36], y=1.95, size=17, min_row=.70)
    d1, d2 = (data['metrics'][f'D{i}'] for i in (1, 2))
    delta1 = (d1['after']['top_prefix_share'] - d1['before']['top_prefix_share']) * 100
    delta2 = (d2['after']['top_prefix_share'] - d2['before']['top_prefix_share']) * 100
    deck.text(stats, f'主前缀目标占比：D1 {delta1:+.2f}pp；D2 {delta2:+.2f}pp', end + .26, .42, size=21)
    deck.text(stats, '主前缀占比 = 最大前缀的目标订单份额；桶大小按同一份额加权。', end + .83, .40, size=18)
    deck.text(stats, '完整方法的 Train 结构变化，不等同于 Query 单项贡献或检索收益。', end + 1.28, .40, size=18)

    extra = []
    if 'tsne.json' in manifest['artifacts']:
        projection = read_json(analysis_dir / 'tsne.json')
        page = deck.page('Query 关联结构：S1/S2 码向量 t-SNE')
        deck.text(page, '五组 D2 改善样例，同一批 POI；固定 seed=42，联合投影，无位置抖动。',
                  1.03, .42, size=19)
        for column, (variant, name) in enumerate((('before', 'PRQ-KMeans（POI-only）'), ('after', 'QG-HRQ'))):
            deck.text(page, name, 1.57, .42, x=.67 + column * 6.18, w=5.80, size=21, align=PP_ALIGN.CENTER)
            page.shapes.add_picture(str(analysis_dir / f'query_tsne_{variant}.png'),
                                    Inches(.67 + column * 6.18), Inches(2.04), width=Inches(5.80), height=Inches(4.25))
        for index, (query, color) in enumerate(zip(projection['queries'], projection['colors'])):
            x = .75 + index * 2.5
            shape = page.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(6.32), Inches(.14), Inches(.14))
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color[1:])
            shape.line.fill.background()
            case = data['projection_cases'][index]
            label = f"{query}\n{case['before']['prefix_count']} → {case['after']['prefix_count']} 个前缀"
            deck.text(page, label, 6.19, .61, x=x + .21, w=2.20, size=16)
        deck.text(page, '同色点合并表示目标共享前缀；点面积表示 POI 数。改善样例，不代表全量。',
                  6.89, .39, size=18)
        extra.append(page)

    order = [*original[:11], tree, stats, *extra, *original[11:]]
    ids = deck.prs.slides._sldIdLst
    lookup = {item.id: item for item in ids}
    ordered_ids = [lookup[s.slide_id] for s in order]
    for item in list(ids): ids.remove(item)
    for item in ordered_ids: ids.append(item)
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.prs.save(output)
    revised = Presentation(output)
    assert len(revised.slides) == 25 + len(extra)
    by_id = {s.slide_id: s for s in revised.slides}
    for old in Presentation(source).slides:
        new = by_id[old.slide_id]
        assert etree.tostring(old._element, method='c14n') == etree.tostring(new._element, method='c14n')
        assert old.has_notes_slide == new.has_notes_slide
        if old.has_notes_slide:
            assert etree.tostring(old.notes_slide._element, method='c14n') == etree.tostring(new.notes_slide._element, method='c14n')
    with ZipFile(source) as before, ZipFile(output) as after:
        assert after.testzip() is None
        for name in before.namelist():
            if name.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/')):
                roots = [etree.fromstring(z.read(name)) for z in (before, after)]
                if name.endswith('.rels'):
                    for root in roots: root[:] = sorted(root, key=lambda e: e.get('Id'))
                assert etree.tostring(roots[0], method='c14n') == etree.tostring(roots[1], method='c14n')
    tables = [s.table for p in revised.slides for s in p.shapes if s.has_table]
    assert len(tables) == 8
    for table in tables:
        for row in list(table.rows)[1:]:
            assert all(str(c.fill.fore_color.rgb) == 'FFFFFF' for c in row.cells)
            assert all(not r.font.bold for c in row.cells for p in c.text_frame.paragraphs for r in p.runs)
    assert digest(source) == source_hash
    record = {'source': str(source), 'source_sha256': source_hash, 'output': str(output),
              'output_sha256': digest(output), 'analysis_manifest_sha256': digest(analysis_dir / 'manifest.json'),
              'slides': len(revised.slides), 'unchanged_source_slides': 23, 'native_tables': 8,
              'masters_layouts_themes_notes_unchanged': True, 'selection': data['selection'], 'contract': data['contract']}
    record_dir = analysis_dir / output.stem
    record_dir.mkdir(exist_ok=False)
    (record_dir / 'ppt_manifest.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'output': str(output), 'slides': len(revised.slides), 'unchanged_source_slides': 23}, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='按上传原稿版式修订 0917 PPT；只改内容，不改母版。')
    parser.add_argument('--output', type=Path, default=OUTPUT, help='新 PPT 路径，禁止覆盖已有文件')
    parser.add_argument('--validate-only', type=Path, help='只读核验已生成文件的原版式与表格合同')
    parser.add_argument('--s3-figures', type=Path, help='仅替换 S3 可视化页，读取已完成图表的 manifest')
    parser.add_argument('--query-analysis', type=Path, help='追加 Query 前缀树和全量统计页，读取已完成分析的 manifest')
    parser.add_argument('--source', type=Path, default=OUTPUT, help='S3 替换模式的输入 PPT，默认已确认的原版式修订稿')
    args = parser.parse_args()
    if args.validate_only:
        validate_output(args.validate_only, digest(SOURCE))
    elif args.query_analysis:
        append_query_visualizations(args.source, args.query_analysis, args.output)
    elif args.s3_figures:
        replace_s3_visualizations(args.source, args.s3_figures, args.output)
    else:
        build(args.output)
