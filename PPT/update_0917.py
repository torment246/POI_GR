#!/usr/bin/env python3
"""Revise the QG-HRQ report using editable objects and frozen evidence."""

from __future__ import annotations

import argparse
from copy import deepcopy
from io import BytesIO
import hashlib
import json
import math
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from build_0907 import Deck, style_frame, wrapped, font

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "PPT/胡丹-0917.pptx"
WORK = ROOT / "qg_prqk/outputs/ppt_0917"
OUTPUT = ROOT / "PPT/胡丹-0917-QG-HRQ完善版.pptx"
FIGURES = ROOT / "qg_prqk/outputs/figures"
SID_ROOT = ROOT / "qg_prqk/outputs/qg_prqk_512x3_v2_1_cat_active"
METHODS = ["A0_POI_ONLY", "A4_GID_PARENT", "A4_NOGID_S1S2_PARENT"]
LABELS = ["A0：POI-only", "QG-HRQ：GID", "QG-HRQ：无 GID"]


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class Report(Deck):
    """Keep the uploaded slide layouts and reuse the 0810 native table style."""

    def __init__(self) -> None:
        self.prs = Presentation(SOURCE)
        self.original = list(self.prs.slides)
        self.title = deepcopy(self.original[1].shapes[0]._element)
        self.layout = self.original[11].slide_layout
        template = Presentation(ROOT / "PPT/胡丹-0810.pptx")
        self.table_props = deepcopy(next(s for s in template.slides[7].shapes if s.has_table).table._tbl.tblPr)
        self.expected_tables: list[tuple[object, list[list[str]]]] = []
        self.sources: dict[str, str] = {}

    def clean(self, slide, title: str, claim: str, source: str = ""):
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        slide.shapes._spTree.insert_element_before(deepcopy(self.title), "p:extLst")
        heading = slide.shapes[0]
        heading.width = Inches(12.7)
        style_frame(heading.text_frame, title, 28, "000000")
        self.text(slide, claim, 1.05, 0.64, size=17, color="1A3A6B")
        if source:
            self.text(slide, source, 6.98, 0.27, size=9.5, color="64748B", w=11.6)
        return slide

    def added(self, title: str, claim: str, source: str = ""):
        return self.clean(self.prs.slides.add_slide(self.layout), title, claim, source)

    def table(self, slide, headers, rows, widths, **kwargs):
        end = super().table(slide, headers, rows, widths, **kwargs)
        self.expected_tables.append((slide, [[str(c) for c in r] for r in [headers, *rows]]))
        return end

    def node(self, slide, label: str, x: float, y: float, w: float, h: float):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = "method_node"
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string("F8FAFC")
        shape.line.color.rgb = RGBColor.from_string("CBD5E1")
        shape.line.width = Pt(0.7)
        style_frame(shape.text_frame, label, 16, "111827", align=PP_ALIGN.CENTER)
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        return shape

    def arrow(self, slide, x1: float, y1: float, x2: float, y2: float):
        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shape.line.color.rgb = RGBColor.from_string("64748B")
        shape.line.width = Pt(1.2)
        end = OxmlElement("a:tailEnd"); end.set("type", "triangle")
        shape.line._get_or_add_ln().append(end)


def motivation(deck: Report) -> None:
    slide = deck.clean(deck.original[1], "为什么要改 SID：分组相似不等于查询可检索",
        "QG-HRQ 聚焦离线索引：让 SID 的粗到细结构与 Query 能提供的监督粒度匹配。",
        "问题定义：本项目方法假设；示意表达不代表真实 Query 的监督标签。")
    nodes = [("POI 内容向量", .65), ("离线量化\nS1 / S2 / S3", 3.78),
             ("Query + 位置 + 历史\n预测目标 SID", 6.91), ("有限 Beam\n检索到目标 POI", 10.04)]
    for label,x in nodes: deck.node(slide,label,x,1.95,2.65,.95)
    for x in [3.3,6.43,9.56]: deck.arrow(slide,x,2.42,x+.44,2.42)
    deck.text(slide,"同一名称不等于同一实体",3.35,.4,size=20,bold=True)
    deck.text(slide,"跨区连锁门店可能文本近似，但目标依赖位置；过早按地理拆分，又可能破坏共享语义。",3.87,.70,size=18)
    deck.text(slide,"一次发单不等于 Query 只指向一个 POI",4.85,.4,size=20,bold=True)
    deck.text(slide,"泛类词可能有多个合理目标；若把每条行为边都压到末层一致，容易把局部选择当成完整意图。",5.37,.75,size=18)

    slide = deck.clean(deck.original[2], "现有研究解决了什么，仍留下什么空间",
        "不是首次引入 Query、类别或地理；差异在监督深度与各层职责。",
        "原文：[1] arxiv.org/abs/2305.05065  [2] /2608.24207  [3] /2605.14434  [4] /2605.03397")
    entries = [
        ("TIGER [1]", "内容向量经 RQ-VAE 得到语义 ID，再学习序列生成。", "没有直接按真实搜索 Query 的目标分布定义逐层监督。"),
        ("PRQ-KMeans [2]", "去全局方向、Top-k 质心精炼、投影残差，改善量化几何。", "实体向量聚类本身不等于 Query 与 POI 的逐层离散对齐。"),
        ("CQ-SID [3]", "首层类别约束与高置信 Query–Item 对比学习；输出语义簇。", "原文未按 Query 歧义设置可监督深度；簇召回与精确 POI 目标也不同。"),
        ("GenPOI [4]", "GeoPE、显式 GID 与近邻约束生成，增强地图空间建模。", "它已建模地理；本文另探索把局部地理限制在末层，并按 Query 粒度对齐。"),
    ]
    for i,(name,done,gap) in enumerate(entries):
        y=1.91+i*1.10
        label=deck.text(slide,name,y,.35,w=2.18,size=19,bold=True)
        label.text_frame.paragraphs[0].runs[0].hyperlink.address='https://arxiv.org/abs/'+['2305.05065','2608.24207','2605.14434','2605.03397'][i]
        deck.text(slide,done,y,.46,x=2.95,w=9.7,size=16)
        deck.text(slide,"本项目切入点："+gap,y+.47,.48,x=2.95,w=9.7,size=16,color="64748B")

    slide=deck.clean(deck.original[3],"QG-HRQ 的假设：只对齐 Query 能确定的前缀",
        "保留 POI 与 Query 各自的连续视图，用可靠行为边约束共享的离散层级。",
        "QG_PRQK_METHOD_SPEC.md；这是设计目标，是否有效由后续静态指标与 SFT 共同检验。")
    rows=[("目标分布分散", "只知道粗类别", "可确定细类别", "可确定精确 POI"),
          ("D0：不监督 SID", "D1：监督 S1", "D2：监督 S1/S2", "D3：监督 S1/S2/S3")]
    for i in range(4):
        x=.65+3.13*i
        deck.node(slide,rows[0][i],x,2.05,2.65,.70)
        deck.arrow(slide,x+1.325,2.80,x+1.325,3.20)
        deck.node(slide,rows[1][i],x,3.25,2.65,.90)
    deck.text(slide,"语义组织：S1/S2",4.68,.4,size=20,bold=True)
    deck.text(slide,"类别作为可分裂的软锚点；同一类别可占多个码字，不硬绑定 CategoryID。",5.17,.7,w=5.7,size=17)
    deck.text(slide,"局部区分：S3",4.68,.4,x=6.94,w=5.7,size=20,bold=True)
    deck.text(slide,"只使用 D3 残差，并在给定父组内加入地理与困难实体约束。",5.17,.7,x=6.94,w=5.7,size=17)
    deck.text(slide,"判断标准：前缀是否更有结构、碰撞是否减少，以及同协议下的目标 POI 检索是否改善。",6.27,.53,size=17)


def update_method_details(deck: Report) -> None:
    # Keep user-authored diagrams and equations; correct only stale contracts.
    slide=deck.original[4]
    style_frame(slide.shapes[18].text_frame,"仅碰撞追加末位 [D]\nGID 版：GID6 + SID3 + [D]",15,"111827")
    slide.shapes[18].height=Inches(.95)
    style_frame(slide.shapes[20].text_frame,
        "D1→S1；D2→S1/S2；D3→S1/S2/S3。\n无 GID 版：SID3+[D]；父组为 (S1,S2)。",15,"111827")
    slide.shapes[20].height=Inches(.80)
    # These examples are explanatory, not data-derived labels.
    deck.text(deck.original[5],"示例只解释监督深度；真实分组由 Train 目标分布决定，不按 Query 字面直接指定。",
        7.0,.25,size=9.5,color="64748B")
    s=deck.original[7]
    deck.text(s,"连续空间各自建模，离散索引由 Query–POI 图对齐；不是强迫两类向量完全重合。",1.08,.64,size=17,color='1A3A6B')
    for label,x,y,w,h in [
        ('POI residual\n内容视图',.65,2.10,2.70,1.02),
        ('Query residual\n查询视图',.65,4.40,2.70,1.02),
        ('POI 质心 U\nPOI assignment',4.02,2.10,2.60,1.02),
        ('Query 质心 V\nQuery assignment',4.02,4.40,2.60,1.02),
        ('共享离散 token k\n图约束促进一致',7.40,3.20,2.32,1.16),
    ]:deck.node(s,label,x,y,w,h)
    # Restore the missing projection target and allow the original text to wrap.
    deck.node(deck.original[8],'选中码向量 c_k\n提取当前方向',10.30,2.55,2.27,.98)
    deck.original[8].shapes[5].height=Inches(.56)
    style_frame(deck.original[8].shapes[5].text_frame,
        '同一粗类可占多个 S1 码字，不硬绑定单一节点。',16,'111827')
    s=deck.original[10]
    deck.node(s,'A0 初始化\nPOI 码本',.65,2.04,1.64,1.12)
    deck.node(s,'分层 Query\n与 POI 图',2.73,2.04,1.64,1.12)
    style_frame(s.shapes[1].text_frame,"先构建 POI-only A0，再以其对应层为初始化，交替更新 POI/Query assignment 与中心。",17,"1A3A6B")
    for index,label in [(3,'更新 POI 分配\n内容 + 图 + 类别'),
                        (5,'更新 Query 分配\n距离 + 图'),
                        (7,'更新双质心 U/V\n与类别分布'),
                        (9,'生成下一层残差\n进入下一层')]:
        style_frame(s.shapes[index].text_frame,label,16,'111827',align=PP_ALIGN.CENTER)
    # Iterate assignments/centroids; residual propagation happens only on exit.
    s.shapes[11].left=Inches(9.76)
    s.shapes[12].width=Inches(9.76-4.78)
    style_frame(s.shapes[14].text_frame,'未收敛：继续迭代',15,'64748B')
    s.shapes[14].left=Inches(6.60)
    s.shapes[14].width=Inches(2.30)
    deck.text(s,'收敛后',1.70,.28,x=10.65,w=1.0,size=11,color='64748B')
    style_frame(s.shapes[18].text_frame,"目标相对改善 <1e-4；POI 变化 <0.1%；Query <0.2%；\n连续两轮满足，或最多 60 轮。",15.5,"111827")


def add_evidence(deck: Report) -> list:
    slides=[]
    category_path=FIGURES/"qg_prqk_sid_category_region_v2/metrics.json"
    category=read_json(category_path)
    slides.append(s:=deck.added("SID 可视化：保留整体布局，改变局部类别边界",
        "同样五类、同样 1,000 个 POI、同一投影；不是把不同样本的散点图放在一起。",
        "冻结图：qg_prqk_sid_category_region_v2；A0 为内部 POI-only 对照，非 TIGER / GenPOI 的可视化。"))
    s.shapes.add_picture(str(FIGURES/"qg_prqk_sid_category_region_v2/sid_category_tsne.png"),Inches(.62),Inches(1.70),width=Inches(12.06))
    deck.text(s,"变化并非“重新分成五个清晰簇”；两版 A4 共用 S1/S2，因此图形相近。二维图仅作定性观察。",6.55,.35,size=15)

    slides.append(s:=deck.added("前缀语义：前两层更纯，不只依赖末层单例",
        "全库 716,245 POI；三层统一使用完整 category_code 的 Micro Purity。",
        "来源：qg_prqk_sid_category_region_v2/metrics.json；D1/D2/D3 在本页仅指前缀深度。"))
    colors=["CBD5E1","1A3A6B","94A3B8"]
    for depth in range(3):
        x=.65+4.16*depth
        deck.text(s,["S1：粗前缀","S1/S2：细前缀","S1/S2/S3：完整 SID"][depth],1.96,.4,x=x,w=4.05,size=19,bold=True)
        for j,method in enumerate(METHODS):
            value=category['methods'][method]['catalog_prefix_metrics'][depth]['fine_category']['poi_weighted_top1_share']
            y=2.62+j*.91
            deck.text(s,LABELS[j],y,.3,x=x,w=3.7,size=15)
            bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y+.38),Inches(2.90*value),Inches(.22))
            bar.name='metric_bar';bar.fill.solid();bar.fill.fore_color.rgb=RGBColor.from_string(colors[j]);bar.line.fill.background()
            deck.text(s,f"{value*100:.2f}%",y+.30,.4,x=x+2.95,w=1.05,size=16)
    deck.text(s,"排除单例后，末层细类别纯度仍由 84.91% 提高到 87.47% / 87.12%。",5.56,.54,size=18,bold=True)
    deck.text(s,"边界：类别参与过 A4 监督；这是结构效果，不是独立泛化证据，也不能单独归因于 Query。",6.23,.52,size=16)

    examples=read_json(FIGURES/"qg_prqk_sid_migration_v1/examples.json")
    case=next(c for c in examples['methods']['A4_GID_PARENT'] if c['case_id'].endswith('d2_improved'))
    slides.append(s:=deck.added("真实成员变化：同一个锚点换了哪些邻居",
        "蓝精灵财务顾问(北京)有限公司：S1/S2 从 [440,291] 变为 [247,291]，两边均为 4 个 POI。",
        "来源：qg_prqk_sid_migration_v1/examples.json；固定 seed 的类别提高组案例，不是按最大收益选例。"))
    deck.node(s,"A0：[440,291]\n保留锚点，其余 3 个移出",.65,1.94,5.45,.82)
    deck.node(s,"QG-HRQ：[247,291]\n保留锚点，新增 3 个成员",7.20,1.94,5.45,.82)
    deck.arrow(s,6.22,2.35,7.05,2.35)
    for x,group in [(.68,'removed'),(7.23,'added')]:
        members=case['groups'][group]['preview']
        assert len(members)==3 and not case['groups'][group]['preview_truncated']
        for j,item in enumerate(members):
            deck.text(s,item['displayname'],3.03+j*.55,.49,x=x,w=5.35,size=16)
    deck.text(s,"同细类别同伴：0%",4.86,.45,w=5.4,size=20,bold=True)
    deck.text(s,"同细类别同伴：100%",4.86,.45,x=7.2,w=5.4,size=20,bold=True)
    deck.text(s,"为什么变化：原投资/金融邻域被重新组织为“生活服务:事务所”标签邻域。",5.65,.5,size=17)
    deck.text(s,"但财务顾问与律师并非同一种需求：此例证明类别锚点改变了成员，不能证明 Query 检索更好。",6.27,.52,size=16)

    geo=next(c for c in examples['methods']['A4_NOGID_S1S2_PARENT'] if c['case_id'].endswith('d3_tied'))
    slides.append(s:=deck.added("局部地理案例：类别不变，同桶成员却变近了",
        "NoGID 的 S3 仍使用连续地理特征；不拼接 GID 不等于完全不使用地理信息。",
        "来源：qg_prqk_sid_migration_v1/examples.json；前后均为非单例；网格指 Geohash5。"))
    deck.text(s,geo['anchor']['displayname']+"（共同锚点，房产小区:楼栋号）",1.98,.5,size=20,bold=True)
    for x,group,prefix in [(.65,'removed',geo['before_prefix']),(7.2,'added',geo['after_prefix'])]:
        peer=geo['groups'][group]['preview'][0]
        deck.node(s,('A0' if group=='removed' else 'QG-HRQ 无 GID')+f"：{prefix}",x,2.75,5.45,.65)
        deck.text(s,"锚点：天通北苑1区22号楼-1单元",3.73,.43,x=x,w=5.5,size=16)
        deck.text(s,"同伴："+peer['displayname'],4.30,.43,x=x,w=5.5,size=16)
        deck.text(s,"网格：wx4gc / "+peer['geohash5'],4.88,.43,x=x,w=5.5,size=16)
    deck.arrow(s,6.22,3.07,7.05,3.07)
    deck.text(s,"同类同伴比例：100% → 100%；同网格同伴比例：0% → 100%。",5.63,.52,size=18,bold=True)
    deck.text(s,"这是局部例子，不是全局结论。仅碰撞桶的区域纯度：A0 81.73%，GID 79.19%，NoGID 82.69%。",6.28,.54,size=16)

    slides.append(s:=deck.added("变化的规模与代价：拆开旧碰撞，也引入新碰撞",
        "比较真实同桶成员，而非 SID 数字本身；统计仅含 S1/S2/S3，不拼 GID 或 Dedup。",
        "来源：QG_PRQK.md“ A0→A4 前缀与成员迁移”；qg_prqk_sid_migration_v1/metrics.json。"))
    end=deck.table(s,["全库对比","S1 同伴改变","S1/S2 同伴改变","完整 SID 同伴改变"],
        [["A0 → QG-HRQ GID","100.00%","85.35%","24.02%"],
         ["A0 → QG-HRQ 无 GID","100.00%","85.35%","25.26%"]],[3.6,2.7,2.8,2.96],y=1.9,size=16,min_row=.60)
    deck.text(s,"A0 碰撞 POI：177,206",end+.34,.4,size=19,bold=True)
    for y,label,removed,added,final in [(4.40,"GID-parent",61587,44948,160567),(5.33,"NoGID",62777,52710,167139)]:
        deck.text(s,label,y,.42,w=2.0,size=18,bold=True)
        deck.node(s,f"旧碰撞 → 单例\n{removed:,}",2.7,y-.05,2.55,.72)
        deck.node(s,f"旧单例 → 碰撞\n{added:,}",6.0,y-.05,2.55,.72)
        deck.arrow(s,8.72,y+.30,9.24,y+.30)
        deck.text(s,f"最终碰撞 {final:,}",y+.08,.43,x=9.4,w=3.05,size=17,bold=True)
    deck.text(s,"结论：主要重组前两层类别邻域，末层是局部调整；碰撞净减少，不是所有 POI 都改善。",6.46,.42,size=16)
    return slides


# Baseline cells are transcribed from the user's slides 13–18, not new experiments.
BASELINES = {
    'fixed10k': [[49.78,73.58,79.79,84.51,68.05,74.77],[49.46,73.27,79.61,84.04,67.69,70.24]],
    'seen_query_unseen_pair': [[16.88,35.65,44.55,53.45,34.38,74.52],[15.72,33.83,42.05,50.63,32.42,69.68]],
    'unseen_query_seen_target': [[43.43,59.47,64.82,69.55,56.65,59.48],[44.28,59.82,64.58,68.54,56.71,48.78]],
    'long_tail_target': [[15.74,24.85,29.72,36.58,25.21,54.38],[20.14,29.63,34.14,38.84,29.01,40.65]],
    'cold_target': [[4.45,6.06,7.16,9.48,6.55,49.19],[8.66,14.06,16.72,19.92,13.94,33.73]],
    'test': [[49.88,73.51,79.81,84.49,68.06,74.68],[50.10,73.39,79.60,84.13,67.98,69.96]],
}
TITLES=['固定 10k Validation','泛化：已见 Query / 未见 Query–POI','泛化：新 Query / 已见目标','泛化：长尾目标','泛化：冷目标','7 月 14 日全量 Test']
CLAIMS=[
    '同为无约束时，GID 版相对 TIGER 的 HR@1 / HR@10 为 +0.53 / +0.68pp。',
    '同为无约束时，GID 版 HR@10 比 TIGER 高 5.60pp；约束解码后的增益需单独看。',
    '新表达、已见 POI：GID 版无约束 HR@1 / HR@10 相对 TIGER 为 +3.50 / +2.59pp。',
    '长尾的 Top-K 覆盖改善更明显；合法路径约束为 GID 版增加 10.84pp HR@10。',
    '无约束下两版仍落后 MMBERT；不能把约束后的优势归因于 SID 单独改善。',
    '606,682 条 Test，均为无约束：GID 版相对 TIGER 的 HR@1 / HR@10 为 +1.08 / +1.02pp。',
]


def result_tables(deck: Report) -> None:
    s=deck.clean(deck.original[11],"SID 构建：区分能力与前缀类别组织",
        "北京 active 716,245 POI；布局均为 512×3；只统计三层 SID，不含 GID / Dedup。",
        "基线：用户原第12页；QG 三版：冻结全库静态指标与 category_region_v2/metrics.json。")
    rows=[['TIGER','RQ-VAE','77.03%','33.78%','152','59.29 / 70.17 / 95.82'],
          ['GNPR','RQ-VAE + diversity','85.77%','22.56%','261','64.06 / 76.04 / 98.82'],
          ['GenPOI','RQ-VAE','75.18%','36.58%','191','55.00 / 68.63 / 93.60'],
          ['MMBERT','RQ-VAE','90.60%','16.22%','30','45.72 / 63.71 / 97.17']]
    static=read_json(SID_ROOT/'p8_nogid_s1s2_parent_comparison_v1/full_342879q_716245p/static_metrics.json')
    cat=read_json(FIGURES/'qg_prqk_sid_category_region_v2/metrics.json')
    for name,label in zip(METHODS,['PRQK (A0)','QG-HRQ (GID)','QG-HRQ (无GID)']):
        path=static['methods'][name]['paths']['s1_s2_s3']; levels=cat['methods'][name]['catalog_prefix_metrics']
        rows.append([label,'PRQ-KMeans',f"{path['distinct_ratio_of_pois']*100:.2f}%",f"{(1-levels[2]['singleton_poi_share'])*100:.2f}%",str(int(path['bucket']['max'])),
            ' / '.join(f"{l['fine_category']['poi_weighted_top1_share']*100:.2f}" for l in levels)])
    end=deck.table(s,['方法','Quantizer','唯一 SID\n比例','碰撞 POI\n比例','最大桶','P1 / P2 / P3\nMicro Purity (%)'],rows,
        [2.25,2.05,1.32,1.32,.86,4.26],y=1.90,size=15,min_row=.47,highlight=5)
    deck.text(s,"唯一 SID 比例 = 不同 SID 数 / POI 数；不是单例占比。A4 的类别纯度与碰撞改善并非所有指标都领先。",end+.20,.54,size=15)
    deck.text(s,'各方法输入表示不同，此表不是仅替换 Quantizer 的消融对比。',6.66,.25,size=10.5,color='64748B')
    for index,(key,title,claim) in enumerate(zip(BASELINES,TITLES,CLAIMS)):
        test=key=='test'
        s=deck.clean(deck.original[12+index],title,claim,
            "基线：用户对应截图；QG：qg_prqk/outputs/eval 下三组正式结果；指标单位 %，固定 epoch 3、Beam=10。")
        rows=[[name,'无约束',*[f'{v:.2f}' for v in values]] for name,values in zip(['TIGER','MMBERT'],BASELINES[key])]
        for constrained in ([False] if test else [False,True]):
            run=('sft_epoch3_full_test_20260714_v1' if test else 'sft_epoch3_constrained_fixed10k_generalization_v1' if constrained else 'sft_epoch3_fixed10k_generalization_v1')
            for variant,label in [('a4_gid_parent','QG-HRQ (GID)'),('a4_nogid','QG-HRQ (无GID)')]:
                p=ROOT/'qg_prqk/outputs/eval'/run/variant/'results'
                p=p/'result.json' if test else p/key/'result.json'
                result=read_json(p)
                if result['status']!='completed': raise ValueError(f'结果未完成：{p}')
                metrics=result['metrics'];assert metrics['sample_count']==(606682 if test else 10000)
                row=[label,'约束' if constrained else '无约束']
                row.extend(f'{metrics[m]*100:.2f}' for m in ['hr@1','hr@3','hr@5','hr@10','ndcg@10','valid_id_rate'])
                rows.append(row);deck.sources[str(p.relative_to(ROOT))]=digest(p)
        end=deck.table(s,['方法','解码','HR@1','HR@3','HR@5','HR@10','NDCG@10','Valid ID'],rows,
            [2.65,1.05,1.22,1.22,1.22,1.22,1.75,1.73],y=1.98,size=16,min_row=.57,highlight=2)
        note=('只比较同解码口径；约束行用于观察解码增量，不与无约束基线混作 SID 消融。' if not test else 'Test 不参与选模；数值保留两位小数，完整精度以原始 result.json 为准。')
        deck.text(s,note,end+.33,.62,size=16)
        if key=='unseen_query_seen_target':
            deck.text(s,'已纠正原截图：NoGID 无约束 Valid ID = 51.07%；53.74% 是 MRR@10。',6.63,.28,size=10.5,color='64748B')
        if key=='cold_target':
            deck.text(s,'已纠正原截图：GID 无约束 HR@3 = 9.38%，不是 9.18%。',6.63,.28,size=10.5,color='64748B')


def summary(deck: Report) -> None:
    s=deck.clean(deck.original[-1],"总结：结构改善已验证，机制收益仍需分开归因",
        "QG-HRQ 完成从 Query 粒度监督、层级量化到 SFT / 检索评测的闭环。",
        "来源：本文静态分析与冻结评测；A0 匹配 SFT 结果未在本次材料中给出。")
    items=[('解决的问题','用 Query 可靠监督深度组织前缀：S1/S2 侧重类别语义，S3 侧重局部实体区分。'),
           ('已有证据','相比 A0，GID 版三层最大桶 208→95、碰撞 POI 净减少 16,639；前两层类别更纯。'),
           ('下游结果','无约束 Test 相对 TIGER：HR@1 +1.08pp、HR@10 +1.02pp；改善存在，但幅度有限。'),
           ('仍需回答','静态 Query→S1 探针 29.96%→21.02% 下降。需结合 A0 同协议 SFT 与消融，区分 Query、类别、Geo 各自贡献。')]
    for i,(head,body) in enumerate(items):
        y=1.95+i*1.13
        deck.text(s,head,y,.40,w=2.05,size=20,bold=True)
        deck.text(s,body,y,.88,x=2.92,w=9.7,size=18)


def build(output: Path) -> None:
    if output.resolve()==SOURCE.resolve() or output.exists():
        raise ValueError('输出必须是尚不存在的新文件，禁止覆盖原稿或已有完善版。')
    source_hash=digest(SOURCE)
    deck=Report()
    assert len(deck.original)==19,'源 PPT 页数改变，请重新核对。'
    motivation(deck);update_method_details(deck)
    evidence=add_evidence(deck)
    result_tables(deck);summary(deck)
    # Place evidence immediately after the existing method section.
    order=[*deck.original[:11],*evidence,*deck.original[11:]]
    ids=deck.prs.slides._sldIdLst
    by_id={item.id:item for item in ids}
    ordered_ids=[by_id[s.slide_id] for s in order]
    for item in list(ids): ids.remove(item)
    for item in ordered_ids: ids.append(item)
    # The inherited colored header makes black titles unreadable; keep its geometry.
    for master in deck.prs.slide_masters:
        for shape in master.shapes:
            if not shape.is_placeholder and shape.top<Inches(.85):
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('FFFFFF' if shape.height>Inches(.1) else 'E2E8F0')
    for i,s in enumerate(deck.prs.slides):
        for shape in s.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        if (shape.is_placeholder and shape.placeholder_format.type in (1,3)) or (run.font.size and run.font.size.pt>=28):
                            run.font.color.rgb=RGBColor(0,0,0)
        if i:
            deck.text(s,f'{i:02d}',7.03,.25,x=12.05,w=.65,size=9.5,color='64748B',align=PP_ALIGN.RIGHT)
    # Assert every result cell is a native table, not a picture of a table.
    for s,expected in deck.expected_tables:
        actual=next(x.table for x in s.shapes if x.has_table)
        assert [[c.text for c in r.cells] for r in actual.rows]==expected
    for s in deck.original[11:18]: assert not any(x.shape_type==13 for x in s.shapes)
    assert not any('BeamRisk' in x.text for s in deck.prs.slides for x in s.shapes if x.has_text_frame)
    output.parent.mkdir(parents=True,exist_ok=True)
    deck.prs.save(output)
    with ZipFile(output) as z: assert z.testzip() is None
    reopened=Presentation(output);assert len(reopened.slides)==24
    assert digest(SOURCE)==source_hash,'原稿发生变化。'
    WORK.mkdir(parents=True,exist_ok=True)
    audit={'source':str(SOURCE),'source_sha256':source_hash,'output':str(output),'output_sha256':digest(output),
        'slides':len(reopened.slides),'native_result_tables':7,'metric_sources':deck.sources,
        'baseline_metrics_source':'user supplied slide screenshots 12-18; values transcribed without filling blank rows',
        'paper_sources':['https://arxiv.org/abs/2305.05065','https://arxiv.org/abs/2608.24207','https://arxiv.org/abs/2605.14434','https://arxiv.org/abs/2605.03397']}
    (WORK/'build_manifest.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output':str(output),'slides':len(reopened.slides),'source_unchanged':True},ensure_ascii=False))


def preview(path: Path, directory: Path) -> None:
    """Render a geometry preview, not a PowerPoint/Office rendering."""
    directory.mkdir(parents=True,exist_ok=True)
    prs=Presentation(path);issues=[];pages=[];scale=144

    def text(draw, frame, box, label):
        x,y,w,h=box
        left=x+frame.margin_left/914400*scale
        top=y+frame.margin_top/914400*scale
        available=w-(frame.margin_left+frame.margin_right)/914400*scale
        lines=[]
        for p in frame.paragraphs:
            sz=next((r.font.size.pt for r in p.runs if r.font.size),16)
            color='111827'
            for r in p.runs:
                try: color=str(r.font.color.rgb);break
                except (TypeError,AttributeError): pass
            for line in wrapped(p.text,max(1,available)/2,sz): lines.append((line,sz,color,p.alignment))
        needed=sum(sz*2*1.2 for _,sz,_,_ in lines)
        if needed>h+5 and any(t for t,_,_,_ in lines):issues.append(f'{label}: estimated {needed:.0f}px > {h:.0f}px')
        if frame.vertical_anchor==MSO_ANCHOR.MIDDLE:top+=max(0,(h-needed)/2)
        for line,sz,color,align in lines:
            ft=font(sz);xx=left
            if align==PP_ALIGN.CENTER:xx+=max(0,(available-ft.getlength(line))/2)
            elif align==PP_ALIGN.RIGHT:xx+=max(0,available-ft.getlength(line))
            draw.text((xx,top),line,fill='#'+color,font=ft,anchor='lt');top+=sz*2*1.2

    def paint(shape, canvas, draw, i, transform=(0,0,1,1)):
        ox,oy,sx,sy=transform
        x=ox+shape.left/914400*scale*sx;y=oy+shape.top/914400*scale*sy
        w=shape.width/914400*scale*sx;h=shape.height/914400*scale*sy
        if shape.shape_type==6:
            tr=shape._element.grpSpPr.xfrm
            nsx=sx*tr.ext.cx/tr.chExt.cx;nsy=sy*tr.ext.cy/tr.chExt.cy
            for child in shape.shapes:paint(child,canvas,draw,i,(x-tr.chOff.x/914400*scale*nsx,y-tr.chOff.y/914400*scale*nsy,nsx,nsy))
        elif shape.shape_type==13:
            im=Image.open(BytesIO(shape.image.blob)).convert('RGBA')
            iw,ih=im.size
            im=im.crop((int(iw*shape.crop_left),int(ih*shape.crop_top),int(iw*(1-shape.crop_right)),int(ih*(1-shape.crop_bottom))))
            im=im.resize((max(1,round(w)),max(1,round(h))),Image.Resampling.LANCZOS)
            canvas.paste(im,(round(x),round(y)),im)
        elif shape.has_table:
            yy=y
            for row in shape.table.rows:
                hh=row.height/914400*scale;xx=x
                for ci,c in enumerate(row.cells):
                    ww=shape.table.columns[ci].width/914400*scale
                    draw.rectangle((xx,yy,xx+ww,yy+hh),fill='#'+str(c.fill.fore_color.rgb),outline='#CBD5E1')
                    text(draw,c.text_frame,(xx,yy,ww,hh),f'page {i} table {ci}');xx+=ww
                yy+=hh
        elif shape._element.tag.endswith('}cxnSp'):
            bx=ox+shape.begin_x/914400*scale*sx;by=oy+shape.begin_y/914400*scale*sy
            ex=ox+shape.end_x/914400*scale*sx;ey=oy+shape.end_y/914400*scale*sy
            draw.line((bx,by,ex,ey),fill='#64748B',width=2)
            if shape._element.xpath('.//a:tailEnd[@type="triangle"]'):
                angle=math.atan2(ey-by,ex-bx);ux,uy=math.cos(angle),math.sin(angle)
                draw.polygon([(ex,ey),(ex-10*ux+4*uy,ey-10*uy-4*ux),(ex-10*ux-4*uy,ey-10*uy+4*ux)],fill='#64748B')
        elif shape.has_text_frame:
            gradients=shape._element.xpath('./p:spPr/a:gradFill/a:gsLst/a:gs/a:srgbClr')
            if gradients:
                start,end=[tuple(bytes.fromhex(e.get('val'))) for e in (gradients[0],gradients[-1])]
                for offset in range(max(1,round(w))):
                    ratio=offset/max(1,w-1)
                    color=tuple(round(a+(b-a)*ratio) for a,b in zip(start,end))
                    draw.line((x+offset,y,x+offset,y+h),fill=color)
            else:
                try:fill='#'+str(shape.fill.fore_color.rgb)
                except (AttributeError,TypeError):fill=None
                if fill:
                    draw.rectangle((x,y,x+w,y+h),fill=fill)
            text(draw,shape.text_frame,(x,y,w,h),f'page {i} {shape.name}')
    for i,s in enumerate(prs.slides,1):
        canvas=Image.new('RGB',(1920,1080),'white');draw=ImageDraw.Draw(canvas)
        if s._element.get('showMasterSp','1')!='0':
            for inherited in (s.slide_layout.slide_master,s.slide_layout):
                for shape in inherited.shapes:
                    if not shape.is_placeholder:paint(shape,canvas,draw,i)
        for shape in s.shapes:paint(shape,canvas,draw,i)
        canvas.save(directory/f'slide-{i:02d}.png')
        pages.append(canvas.resize((640,360)))
    for start in range(0,len(pages),6):
        sheet=Image.new('RGB',(1920,800),'#E2E8F0');draw=ImageDraw.Draw(sheet)
        for j,page in enumerate(pages[start:start+6]):
            xx=(j%3)*640;yy=(j//3)*400
            sheet.paste(page,(xx,yy+32));draw.text((xx+12,yy+4),str(start+j+1),font=font(10),fill='black')
        sheet.save(directory/f'contact-{start//6+1:02d}.png')
    print(json.dumps({'preview':'geometry_only_not_office','estimated_text_issues':issues},ensure_ascii=False))


def inspect() -> None:
    """Extract supplied screenshot assets without changing the source deck."""
    target = WORK / "source_images"
    target.mkdir(parents=True, exist_ok=True)
    prs = Presentation(SOURCE)
    for i, slide in enumerate(prs.slides, 1):
        if i < 12:
            continue
        for j, shape in enumerate(slide.shapes):
            if shape.shape_type == 13:
                path = target / f"slide_{i:02d}_{j}.{shape.image.ext}"
                path.write_bytes(shape.image.blob)
                print(path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="补齐 0917 QG-HRQ 汇报；原始 PPT 保持不变。")
    parser.add_argument("--inspect", action="store_true", help="只提取结果截图供核对")
    parser.add_argument("--output", type=Path, default=OUTPUT, help="另存的完善版 PPTX；拒绝覆盖已有文件")
    parser.add_argument("--preview", type=Path, help="只对给定 PPTX 生成非 Office 几何预览")
    args = parser.parse_args()
    if args.preview:
        preview(args.preview,WORK/'preview')
    elif args.inspect:
        inspect()
    else:
        build(args.output)
