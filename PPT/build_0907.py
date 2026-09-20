#!/usr/bin/env python3
"""Build the September phase report from the validated 0810 template and local evidence."""

from __future__ import annotations

import argparse
from copy import deepcopy
from functools import lru_cache
from io import BytesIO
import json
import math
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "PPT/胡丹-0810.pptx"
FONT_PATH = ROOT / "outputs/ppt_0907/assets/NotoSansCJKsc-Regular.otf"
FONT = "Microsoft YaHei"
NAVY, INK, MUTED = "1A3A6B", "111827", "64748B"
HEADER, LIGHT, BLUE = "E2E8F0", "F8FAFC", "EAF2FF"
W = 12.06


def load(path: str) -> dict:
    """Load a frozen result, rejecting failed or partial terminal artifacts."""
    data = json.loads((ROOT / path).read_text())
    if "status" in data and data["status"] not in ("completed", "passed", "ready"):
        raise ValueError(f"非完成态来源：{path}: {data['status']}")
    return data


def pct(value: float, digits: int = 2) -> str:
    return f"{value * 100:.{digits}f}"


def scores(metrics: dict, digits: int = 2) -> list[str]:
    return [pct(metrics[key], digits) for key in ("hr@1", "hr@10", "ndcg@10")]


@lru_cache(maxsize=128)
def font(size: float) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_PATH), max(1, round(size * 2)))


def wrapped(text: str, width_pt: float, size: float) -> list[str]:
    """Conservative CJK wrapping for editable shape layout and preview checks."""
    ft = font(size)
    lines: list[str] = []
    for raw in text.split("\n"):
        current = ""
        for character in raw:
            if current and ft.getlength(current + character) > width_pt * 2 / 1.035:
                lines.append(current)
                current = ""
            current += character
        lines.append(current)
    return lines


def style_frame(frame, text: str, size: float = 16, color: str = INK,
                bold: bool = False, align=PP_ALIGN.LEFT) -> None:
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    for i, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = 1.2
        paragraph.space_after = Pt(4)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        # Explicit East Asian font avoids theme-dependent fallback in PowerPoint.
        ea = OxmlElement("a:ea")
        ea.set("typeface", "微软雅黑")
        run._r.get_or_add_rPr().append(ea)


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation(TEMPLATE)
        table_slide = self.prs.slides[7]
        self.layout = table_slide.slide_layout
        self.title = deepcopy(next(s for s in table_slide.shapes if s.is_placeholder)._element)
        self.page = deepcopy(next(s for s in table_slide.shapes if s.has_text_frame and s.text == "07")._element)
        self.table_props = deepcopy(next(s for s in table_slide.shapes if s.has_table).table._tbl.tblPr)
        self.cover_shapes = [deepcopy(s._element) for s in self.prs.slides[0].shapes]
        self.cover_layout = self.prs.slides[0].slide_layout
        self.source_paths: set[str] = set()
        for sid in list(self.prs.slides._sldIdLst):
            self.prs.part.drop_rel(sid.rId)
            self.prs.slides._sldIdLst.remove(sid)

    def text(self, slide, text: str, y: float, h: float, *, x: float = 0.62,
             w: float = W, size: float = 16, color: str = INK,
             bold: bool = False, align=PP_ALIGN.LEFT):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        style_frame(shape.text_frame, text, size, color, bold, align)
        return shape

    def slide(self, title: str, claim: str, source: str = ""):
        slide = self.prs.slides.add_slide(self.layout)
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        for element in (self.title, self.page):
            slide.shapes._spTree.insert_element_before(deepcopy(element), "p:extLst")
        heading, page = slide.shapes[0], slide.shapes[1]
        heading.width = Inches(12.5)
        style_frame(heading.text_frame, title, 28, "000000")
        style_frame(page.text_frame, f"{len(self.prs.slides)-1:02d}", 9, MUTED, align=PP_ALIGN.CENTER)
        self.text(slide, claim, 1.04, 0.70, size=17, color=NAVY)
        if source:
            self.text(slide, "来源：" + source, 6.94, 0.30, w=11.30, size=9.5, color=MUTED)
        return slide

    def table(self, slide, headers: list[str], rows: list[list[str]], widths: list[float],
              y: float = 1.93, highlight: int | None = None, size: float = 16,
              min_row: float = 0.53) -> float:
        assert abs(sum(widths) - W) < 0.01
        all_rows = [headers] + rows
        heights = []
        for row in all_rows:
            max_lines = max(len(wrapped(str(t), (cw - 0.18)*72, size)) for t, cw in zip(row, widths))
            heights.append(max(min_row, (max_lines * size * 1.2 + 12) / 72))
        if y + sum(heights) > 6.52:
            raise ValueError(f"表格溢出，需要拆页：{slide.shapes[0].text}: {y+sum(heights):.2f}")
        shape = slide.shapes.add_table(len(all_rows), len(headers), Inches(0.62), Inches(y),
                                      Inches(W), Inches(sum(heights)))
        table = shape.table
        table._tbl.replace(table._tbl.tblPr, deepcopy(self.table_props))
        for j, width in enumerate(widths):
            table.columns[j].width = Inches(width)
        for i, row in enumerate(all_rows):
            table.rows[i].height = Inches(heights[i])
            for j, value in enumerate(row):
                cell = table.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(0.09)
                cell.margin_top = cell.margin_bottom = Inches(0.06)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                bg = HEADER if i == 0 else BLUE if i - 1 == highlight else LIGHT if i % 2 == 0 else "FFFFFF"
                cell.fill.fore_color.rgb = RGBColor.from_string(bg)
                style_frame(cell.text_frame, str(value), size, NAVY if i == 0 else INK,
                            bold=(i == 0 or i - 1 == highlight))
                cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.space_after = Pt(0)
                tc_pr = cell._tc.get_or_add_tcPr()
                for edge in ("lnL", "lnR", "lnT", "lnB"):
                    element = OxmlElement(f"a:{edge}")
                    element.set("w", "6350")
                    fill = OxmlElement("a:solidFill")
                    color = OxmlElement("a:srgbClr")
                    color.set("val", "CBD5E1")
                    fill.append(color)
                    element.append(fill)
                    tc_pr.append(element)
        return y + sum(heights)

    def takeaway(self, slide, text: str, y: float) -> None:
        lines = len(wrapped(text, (W - 0.06) * 72, 16))
        h = (lines * 19.2 + 8) / 72
        if y + h > 6.85:
            raise ValueError(f"结论溢出：{slide.shapes[0].text}")
        self.text(slide, text, y, h, size=16)

    def cover(self) -> None:
        slide = self.prs.slides.add_slide(self.cover_layout)
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        for element in self.cover_shapes:
            slide.shapes._spTree.insert_element_before(deepcopy(element), "p:extLst")
        self.text(slide, "阶段实验与方法进展", 4.03, 0.48, x=2.4, w=8.53,
                  size=22, color=NAVY, align=PP_ALIGN.CENTER)
        self.text(slide, "2026.08.25—09.07  |  胡丹", 4.64, 0.40, x=2.4, w=8.53,
                  size=16, color=MUTED, align=PP_ALIGN.CENTER)


def build_content(deck: Deck) -> None:
    """Author only the new phase; historical tables are context, not new claims."""
    tiger_path = "outputs/eval/tiger_bge_m3_1024x3_history10_query_gid_v1_gpu4_6000d_e3/runs/valid_checkpoint-16713_beam10_subset10000/result.json"
    tiger = load(tiger_path)["metrics"]
    tiger_legal = load("outputs/eval/legal_path_fixed10k_v1/tiger_e3/runs/valid_checkpoint-16713_beam10_legalpath_subset10000/result.json")["metrics"]
    bucket = load("outputs/bucket_rerank/EXP-20260828-01_tiger_fixed10k_v1/metrics.json")
    bucket_e4 = load("outputs/bucket_rerank/EXP-20260828-02_e4_512x1024x2048_fixed10k_v1/metrics.json")
    br = load("beamrisk_sft/outputs/eval/main_v1/final_eval_summary.json")
    qg = load("qg_prqk/outputs/qg_prqk_1024x3_v2_1_cat_active/query_adapter_exact_full/comparison/three_view_comparison.json")
    assert qg["candidate_catalog_rows"] == 716245 and qg["selected_epoch"] == 2
    assert br["fixed_validation"]["output_rows"] == 10000
    deck.cover()

    s = deck.slide("本阶段进展", "从替换静态 SID，转向前缀可生成性、训练目标与候选目录的系统验证。",
                   "各方法实验文档与正式 result / manifest；截至 2026-09-07")
    deck.table(s, ["实验线", "本阶段新增工作", "当前结论 / 状态"], [
        ["桶内重排", "固定生成器，展开语义桶后重排 POI", "TIGER / E4 两轮均未超过原 C"],
        ["TIGER-Joint", "从头联合训练 RQ-VAE 与 Qwen", "三轮与双解码完成；第一版停止"],
        ["GID 顺序", "同一 SID 比较 GID-first / SID-first", "四项正式评测完成；均未超过 TIGER"],
        ["BeamRisk-SFT", "真实 Beam 边界自挖掘与风险损失", "三轮完成；负结果与问题已记录"],
        ["active 新基线", "71.6 万 POI，BGE / MMBERT 配对", "512×3 SID 与 SFT 数据完成；等待训练"],
        ["QG-PRQK", "分层 Query + 类别软锚点 + 局部地理", "Exact Adapter 完成；后续 SID 待评审"],
    ], [2.05, 5.1, 4.91], min_row=0.60)

    s = deck.slide("新增诊断：MMBERT 主要损失在粗层", "更少的碰撞没有带来更高目标桶覆盖，继续扩展后缀不能解决主要 miss。",
                   "VECTOR_MODEL_OPTIMIZATION.md，EXP-20260826-01；固定 Validation 10k")
    end = deck.table(s, ["诊断口径", "TIGER", "旧版 MMBERT"], [
        ["正确前缀下 S1 Top-1", "75.06%", "69.06%"],
        ["正确前缀下 S1-S3 累计 Top-1", "53.98%", "50.17%"],
        ["自由生成：唯一 Bucket HR@10", "88.06%", "84.19%"],
        ["同次诊断：精确 POI HR@10", "87.16%", "83.83%"],
        ["只修复桶内排序的 HR@10 理想空间", "+0.90pp", "+0.36pp"],
    ], [7.06, 2.5, 2.5])
    deck.takeaway(s, "MMBERT 的 1,617 条精确 Top-10 miss 中，1,581 条没有目标桶。这里使用同次诊断结果；不混用旧跨硬件评测的 83.84%。", end + 0.25)

    s = deck.slide("Bucket 重排：方法", "先验证现有轻量信号能否兑现 Bucket 覆盖上界，不新增 SFT。",
                   "BUCKET_RERANK.md；src/poi_gr/methods/bucket_rerank/")
    end = deck.table(s, ["步骤", "具体做法"], [
        ["候选生成", "复用 epoch 3 无约束 Beam=10；抽取 S1/S2/S3，按首次出现去重合法桶"],
        ["目录展开", "从冻结 mapping 取出桶内全部 POI；不改变 SID、Qwen 或原始 Beam"],
        ["三类信号", "Train 热度 log1p(count)；Query 与名称/地址的词法分数；冻结向量点积"],
        ["排序对照", "保持桶顺序、仅桶内排序；或全局排序 / RRF 融合，共八种变体"],
        ["评测目标", "最终精确 POI Top-10；与原完整 [S1,S2,S3,C] 候选比较"],
    ], [2.05, 10.01])
    deck.takeaway(s, "E4 组只将 POI 语义分数换为 E4 向量。两组均不使用地理特征、不训练排序器、不在 Validation 调参。", end + 0.25)

    s = deck.slide("Bucket 重排：结果", "两轮八变体均未通过门槛；无语义编号 C 已被生成模型学成条件排序信号。",
                   "outputs/bucket_rerank/EXP-20260828-01、02 的 metrics.json；单位 %")
    rows = [["TIGER 原 C"] + scores(bucket["rankings"]["original_c_token"]["exact"])]
    for label, key in [("TIGER 桶内热度", "bucket_then_popularity"), ("TIGER 桶内词法", "bucket_then_lexical"), ("TIGER 桶内 BGE", "bucket_then_bge")]:
        rows.append([label] + scores(bucket["rankings"][key]["exact"]))
    rows += [["E4 后移原 C"] + scores(bucket_e4["rankings"]["original_c_token"]["exact"]),
             ["E4 后移桶内热度"] + scores(bucket_e4["rankings"]["bucket_then_popularity"]["exact"])]
    end = deck.table(s, ["方法", "HR@1", "HR@10", "NDCG@10"], rows, [5.16, 2.3, 2.3, 2.3], highlight=0)
    deck.takeaway(s, "TIGER / E4 展开池目标覆盖为 88.06% / 88.45%，不等于最终 Top-10。停止直接替换 C；保留候选展开设施，候选补全尚未实施。", end + 0.22)

    s = deck.slide("TIGER-Joint：让 SID 与生成器共同更新", "用连续对齐损失连接 Qwen 与 RQ-VAE；离散 hard SID 本身不能传递生成梯度。",
                   "TIGER_JOINT.md §4–5；src/poi_gr/methods/tiger_joint/")
    end = deck.table(s, ["模块 / 目标", "具体实现与梯度作用"], [
        ["从头初始化", "冻结 BGE POI 向量；fresh KMeans + 1024×3 RQ-VAE；Qwen 基座重新扩词"],
        ["动态训练标签", "原始 POI 行号进入 batch；当前 RQ-VAE 为历史与目标实时生成三层 hard SID"],
        ["L_gen：生成 CE", "预测目标 SID；梯度只更新 Qwen，不穿过 hard argmin"],
        ["L_align：InfoNCE", "当前请求 hidden → 256D；对齐目标量化表示，连接 Qwen、投影与 RQ encoder"],
        ["L_rq：目录重建", "独立均匀目录批覆盖全部 POI；更新 encoder、decoder 与 codebook"],
    ], [2.85, 9.21])
    deck.takeaway(s, "总目标 L = L_gen + 0.1 L_align + L_rq。每次更新后 SID 可能变化；epoch 3 冻结并导出目录，只评三级 Bucket，不追加 C。", end + 0.22)

    joint_u = load("outputs/eval/tiger_joint/EXP-20260829-01_epoch3_bucket_fixed10k/result.json")["bucket_metrics"]
    joint_c = load("outputs/eval/tiger_joint/EXP-20260830-01_epoch3_constrained_bucket_fixed10k/result.json")["bucket_metrics"]
    s = deck.slide("TIGER-Joint：合法性修复未扭转结果", "完成四卡三轮、44,454 step；首层码使用减少，正确终态桶的概率排序仍明显不足。",
                   "outputs/eval/tiger_joint/EXP-20260829-01、20260830-01；Bucket 指标，单位 %")
    end = deck.table(s, ["模型 / 解码", "Bucket HR@1", "Bucket HR@10", "候选可展开率"], [
        ["TIGER / 无约束参考", "55.53", "88.06", "76.13"],
        ["Joint / 无约束", pct(joint_u["unique_bucket_hr@1"]), pct(joint_u["unique_bucket_hr@10"]), pct(joint_u["expandable_candidate_rate"])],
        ["Joint / 全目录约束", pct(joint_c["unique_bucket_hr@1"]), pct(joint_c["unique_bucket_hr@10"]), "100.00"],
    ], [4.26, 2.6, 2.6, 2.6])
    deck.takeaway(s, "全量三层 SID 不同组合占比 81.93%，S1 只使用 425/1,024 个码。约束使 Joint 的 HR@10 增加 9.67pp，但 HR@1 仅增加 0.15pp。", end + 0.28)
    deck.takeaway(s, "当前关闭逐 batch 动态 hard-label 第一版。标签变化是可能因素，尚未证明唯一因果；约束与无约束应分开解读，不能把 Bucket 命中当精确 POI 命中。", end + 1.10)

    s = deck.slide("GID 前后顺序：成对控制怎么做", "冻结同一份 TIGER SID3 和地理映射，仅改变历史与目标 PID 的生成顺序。",
                   "TIGER.md，EXP-20260830-02/03、20260831-01；两版共享初始化")
    end = deck.table(s, ["部分", "GID-first", "SID-first"], [
        ["目标 / 历史 POI 标识", "G1…G6 → S1 → S2 → S3 → [D]", "S1 → S2 → S3 → G1…G6 → [D]"],
        ["生成假设", "先选区域，再在区域内找语义实体", "先选语义，再用地理定位具体实体"],
        ["共享部分", "同一 POI、SID3、GID6、局部 D", "与左侧完全一致"],
        ["训练控制", "同一扩词初始化、样本、三轮 SFT", "仅 identifier 顺序变化"],
        ["输入请求", "当前 Query + 用户 GID + 历史10", "当前 Query + 用户 GID + 历史10"],
    ], [2.56, 4.75, 4.75])
    deck.takeaway(s, "SID3 不同组合占比 71.68%，加 GID6 后为 82.76%，局部 D 后 100% 唯一；两版目标均为 9/10 个编码 token。", end + 0.22)

    s = deck.slide("GID 前后顺序：正式双解码结果", "GID-first 优于 SID-first，但两版均未超过 TIGER；更低 token loss 不能替代生成评测。",
                   "outputs/eval/pid_order_fixed10k_v1/*/runs/*fixed10k/result.json；epoch 3，单位 %")
    rows = [["TIGER", "无约束"] + scores(tiger), ["TIGER", "全目录约束"] + scores(tiger_legal)]
    for order, label in [("gid_sid", "GID-first"), ("sid_gid", "SID-first")]:
        for mode, mode_label in [("unconstrained", "无约束"), ("constrained", "全目录约束")]:
            result = load(f"outputs/eval/pid_order_fixed10k_v1/{order}_epoch3/runs/{mode}_epoch3_beam10_fixed10k/result.json")
            assert result["metrics"]["sample_count"] == 10000
            rows.append([label, mode_label] + scores(result["metrics"]))
    end = deck.table(s, ["方法", "解码", "HR@1", "HR@10", "NDCG@10"], rows, [2.46, 3.0, 2.2, 2.2, 2.2])
    deck.takeaway(s, "两种顺序是严格成对实验。与旧 TIGER 相比还存在目标长度、cutoff 512/1024 和 packed step 差异，因此不能将全部回退单独归因为 GID。", end + 0.22)

    s = deck.slide("BeamRisk-SFT：从真实 Beam 找风险边界", "SID 不变、从相同初始模型重训；辅助目标直接关注进不了 Top-10 和排不到第一。",
                   "beamrisk_sft/README.md §6；tracing.py、mining.py、loss.py")
    end = deck.table(s, ["当前模型状态", "正负样本如何取", "辅助目标"], [
        ["目标首次掉出 Beam", "首次失败深度的 gold 前缀\n对比当层实际第 10 名边界前缀", "Survive：提高目标生存分数"],
        ["目标进入第 2–10 名", "完整 gold SID\n对比当前第 1 名错误 SID", "Rank：改善最终排序"],
        ["目标已经排第 1 名", "不构造额外风险 pair", "只保留标准 CE"],
    ], [3.0, 5.06, 4.0], min_row=0.72)
    deck.takeaway(s, "Score(path) 为逐 token log 概率之和。实际 v1 损失：L_risk = softplus(Score_negative − Score_gold)，每条请求只取一种风险。", end + 0.30)
    deck.takeaway(s, "设计意图是跨越真实 Beam 边界，但 softplus 不会在刚越过边界后停止梯度；这是本轮结果暴露出的关键问题。", end + 1.10)

    s = deck.slide("BeamRisk-SFT：全量 CE + 独立风险流", "500k 只是 Train-only 挖掘候选池；每轮主 CE 仍遍历完整训练集。",
                   "beamrisk_sft/configs/main_v1.yaml；README §7–8；final_eval_summary.json")
    end = deck.table(s, ["阶段", "如何运行", "数据 / 状态"], [
        ["Epoch 1", "从扩词基座开始普通 CE", "7,586,410 原始 Train；2,851,853 packed rows"],
        ["自挖掘 1", "用自己的 epoch-1 模型跑 Beam=10", "Train-only 候选池；选择 100k 风险 pair"],
        ["Epoch 2", "完整 CE + 每 4 optimizer step 注入风险梯度", "weight=0.2；每次全局 128 pair"],
        ["自挖掘 2 + Epoch 3", "刷新风险 pair，继续同一训练轨迹", "恢复 optimizer / scheduler / 四 rank RNG"],
        ["最终评测", "只评 epoch 3，固定同业务键 10k", "无约束与全目录合法路径约束各一项"],
    ], [2.25, 5.15, 4.66])
    deck.takeaway(s, "三轮共 16,713 step；不加载已训练 TIGER checkpoint，不把风险池替代主训练数据，也不额外追加第四轮。", end + 0.25)

    s = deck.slide("BeamRisk-SFT：结果未超过 TIGER", "两种解码均回退；Epoch 1 基本对齐，风险启用后主 CE 与召回开始落后。",
                   "beamrisk_sft/outputs/eval/main_v1/final_eval_summary.json；单位 %")
    end = deck.table(s, ["方法 / 解码", "HR@1", "HR@10", "NDCG@10"], [
        ["TIGER / 无约束"] + scores(tiger),
        ["BeamRisk / 无约束"] + scores(br["unconstrained"]["metrics"]),
        ["TIGER / 全目录约束"] + scores(tiger_legal),
        ["BeamRisk / 全目录约束"] + scores(br["legal_path_constrained"]["metrics"]),
    ], [5.16, 2.3, 2.3, 2.3])
    deck.takeaway(s, "无约束 HR@1 / HR@10 / NDCG@10 分别下降 1.26 / 0.63 / 0.91pp。已有逐请求配对区间均为负；不是通过合法性约束就能解决的问题。", end + 0.26)
    deck.takeaway(s, "Validation CE：TIGER 0.3848 → 0.3058 → 0.2965；BeamRisk 0.3852 → 0.3167 → 0.3040。普通日志 loss 只记录主 CE，不含额外风险 backward。", end + 1.05)

    s = deck.slide("BeamRisk-SFT：失败证据与冻结决策", "风险梯度强度、样本定义与单点击假负例共同值得排查；不以降低风险 loss 判定成功。",
                   "beamrisk_sft/BEAMRISK_V1_RESULT_ANALYSIS.md；v2 仅记录，未实施")
    end = deck.table(s, ["证据", "本轮观察", "对应问题"], [
        ["梯度裁剪", "Epoch 2 / 3 已记录风险步范数中位数 3.05 / 2.26；均大于 1", "合并裁剪也压缩主 CE 梯度"],
        ["跨桶排序", "Epoch 2 的 final-rank pair 中 91% 首次差异位于 S1–S3", "并非只在桶内优化 C"],
        ["边界过推", "Epoch 3 平均风险 margin 约 −15.16", "softplus 持续追求更大间隔"],
        ["关系型假负例", "北京南站与南进站口互为负例；严格重复过滤未保护", "单点击不代表其他相关 POI 必错"],
    ], [2.05, 6.10, 3.91])
    deck.takeaway(s, "保留 v1 产物，暂不修改或重跑。v2 候选：首剪枝优先、越界后停止的损失、独立梯度预算、同桶排序与关系保护；均尚未验证。", end + 0.26)

    add_active_slides(deck)
    add_qg_slides(deck, qg)

    s = deck.slide("阶段总结与下一步", "最终仍以生成式 HR / NDCG 判断方法；当前优先拿到 active 两组配对结果。",
                   "本轮只制作汇报；未改实验配置、未启动训练或 QG 下一阶段")
    end = deck.table(s, ["方向", "已得到的结论", "下一步边界"], [
        ["桶内重排 / Joint", "简单替换 C 与动态 hard-label v1 均失败", "保留基础设施，不继续扩旧变体"],
        ["GID / BeamRisk", "更长目标或风险 loss 改善不保证召回", "不把低 loss / 高合法率当效果结论"],
        ["active TIGER / MMBERT", "同一 716,245 POI、512×3 和 SFT 协议", "等待训练与固定 10k + 四类泛化结果"],
        ["QG-PRQK", "D3 Adapter 内部验证正向；SID 尚未构建", "P3A-FULL 审核后才推进 P4–P8"],
        ["当前位置输入", "用户 GID 是输入条件，与输出 SID 独立", "当前不要求给目标 SID 再拼 GID"],
    ], [2.76, 5.15, 4.15])
    deck.takeaway(s, "新旧候选库、不同输出结构和不同评测集合分开报告。QG 的 Adapter 增益与 active 的静态 SID 增益，目前都不能表述为超过 TIGER。", end + 0.23)


def add_active_slides(deck: Deck) -> None:
    """Separate the new candidate universe from historical full-catalog conclusions."""
    s = deck.slide("数据范围修订：活跃闭集 POI", "不只保留 Train 目标，而是保留两周全部目标与历史中出现的 POI。",
                   "data/beijing_poi_active_order14d_history10_20260715_json/manifest.json")
    end = deck.table(s, ["集合", "唯一 POI 数", "定义"], [
        ["北京原全目录", "2,337,178", "2026-07-15 在线 POI 快照"],
        ["两周全部目标 A", "520,333", "07-01—14，包含 Train / Valid / Test"],
        ["历史 POI B", "607,456", "用户 04-01—06-30 的最近 10 条历史"],
        ["交集 A ∩ B", "411,544", "两种数据中均出现"],
        ["新目录 A ∪ B", "716,245", "按原 POI 主表行序抽取完整属性"],
    ], [3.65, 2.36, 6.05], highlight=4)
    deck.takeaway(s, "候选规模减少约 69.35%，行为样本与时间切分不变。Train 目标仍为 491,213 个，占 active 约 68.58%，并非新目录每个 POI 都有目标监督。", end + 0.23)

    s = deck.slide("active 基线：同一目录重建两套 SID", "仅替换向量模型，量化、去重及后续 SFT 协议保持配对。",
                   "TIGER.md EXP-20260904-02；VECTOR_MODEL_OPTIMIZATION.md EXP-20260904-03")
    end = deck.table(s, ["阶段", "active-BGE", "active-MMBERT"], [
        ["POI 输入", "716,245 条名称 / 地址 / 别名", "同目录、同行序、同文本"],
        ["冻结向量", "BGE-M3，1,024 维", "线上 checkpoint；mean pooling\n训练好的 768→128 投影 + L2"],
        ["量化模型", "RQ-VAE：1024→512→256", "RQ-VAE：128→512→256"],
        ["码本与训练", "512×512×512，20 epoch", "相同容量、轮数与 seed=42"],
        ["初始化与去重", "全部 716,245 向量初始化；\n桶内按 POI ID 稳定分配 C", "相同初始化范围与 C 分配规则"],
    ], [2.06, 5.0, 5.0])
    deck.takeaway(s, "这是基于两周目标与历史身份定义的活跃闭集，不是未来开放目录评测。量化不使用评测 Query–POI 监督，但缩库本身改变候选协议。", end + 0.24)

    roots = ["outputs/sid/tiger/active_716k_bge_m3/TIGER-ACTIVE716K-BGE-M3-512x3-FULLINIT",
             "outputs/sid/tiger/active_716k_mmbert_recall_128/TIGER-ACTIVE716K-MMBERT-RECALL-128-512x3-FULLINIT"]
    sid = [load(r + "/evaluations/epoch_20/metrics.json") for r in roots]
    ids = [load(r + "/tiger_ids/epoch_20/metrics.json") for r in roots]
    assert all(d["basic"]["poi_count"] == 716245 for d in sid)
    assert all(d["tiger_id_unique_ratio"] == 1 for d in ids)
    s = deck.slide("active SID：MMBERT 静态结构更好", "同一 512×3 容量下，MMBERT 减少碰撞并保持全部码字活跃；尚无新 SFT 指标。",
                   "active 两组 epoch_20/evaluations 与 tiger_ids 的 metrics.json")
    rows = [
        ["不同三层 SID / POI 数"] + [pct(d["basic"]["distinct_sid_ratio"], 4)+"%" for d in sid],
        ["处在碰撞桶的 POI 占比"] + [pct(d["basic"]["colliding_poi_ratio"], 4)+"%" for d in sid],
        ["最大三层桶 / 所需 C token 数"] + [str(d["required_collision_token_count"]) for d in ids],
        ["最终 S1 / S2 / S3 活跃码数"] + [" / ".join(str(sum(n > 0 for n in l["token_counts"])) for l in d["layers"]) for d in sid],
        ["加 C 后四层标识唯一率"] + ["100.00%", "100.00%"],
        ["正式 SFT 结果"] + ["等待训练", "等待训练"],
    ]
    end = deck.table(s, ["指标", "active-BGE", "active-MMBERT"], rows, [6.06, 3.0, 3.0])
    deck.takeaway(s, "“不同 SID / POI 数”不等于单例 POI 占比。以上均为 716,245 条全目录导出指标，不使用小规模 reconstruction monitor 代替。", end + 0.22)

    s = deck.slide("active SFT：两组输入与六项评测已就绪", "TIGER 与 MMBERT 当前都等待训练；训练结束后自动评 epoch 3。",
                   "TIGER.md EXP-20260905-01；VECTOR_MODEL_OPTIMIZATION.md EXP-20260906-01")
    end = deck.table(s, ["项目", "共同协议 / 验收状态"], [
        ["原始样本", "Train / Valid / Test = 7,586,410 / 597,421 / 606,682；只替换 identifier"],
        ["序列安全", "cutoff=1024；Train+Valid 最长 953；超长与目标截断均为 0"],
        ["正式缓存", "packed Train / Valid = 1,296,883 / 96,949；两组均已构建完成"],
        ["训练配置", "Qwen3-0.6B 全参数；4×6000D；batch 8 × 累积16 × 4卡 = 512；3 epoch"],
        ["固定随机 10k", "同业务键 Validation；无约束 Beam=10 + 全目录合法路径约束"],
        ["四类泛化 10k", "已见 Query / 新 Pair；新 Query / 已见 POI；长尾目标；冷目标；均无约束"],
    ], [2.36, 9.70])
    deck.takeaway(s, "已通过全链路 dry-run，不代表已经完成平台实训。下一结论只比较这两组同目录结果，不能把缩库、容量和编码器的影响混为单一收益。", end + 0.22)


def add_qg_slides(deck: Deck, comparison: dict) -> None:
    """Mark prospective tokenizer mechanisms separately from completed Adapter results."""
    src = "qg_prqk/docs/methods/QG_PRQK_METHOD_SPEC.md；canonical v2.1-CAT"
    s = deck.slide("QG-PRQK：最新 SID 方法主线", "按 Query 可确定的语义粒度分层监督 SID，而非把全部 Query 混入 POI 向量。", src)
    end = deck.table(s, ["已有问题", "QG-PRQK 的设计回应", "目前状态"], [
        ["Query 不一定指向具体 POI", "按粗类 / 细类 / 精确实体分 D1/D2/D3，只监督相应层", "分层统计已完成"],
        ["短 Query 与长 POI 向量有差异", "POI 与 Query 独立视图、独立质心，共享离散 token", "后续量化设计"],
        ["粗层可生成性容易退化", "S1 粗类别、S2 条件细类别软锚点，避免硬绑类别 ID", "后续量化设计"],
        ["近距离相似实体难区分", "只在 S3 引入局部相对地理与困难实体软约束", "后续量化设计"],
    ], [3.0, 6.1, 2.96])
    deck.takeaway(s, "与 E4 区别：E4 直接改变每个 POI 向量；QG 保持 POI/Query 双视图，计划在离散分配中对齐。与 Joint 区别：先离线构建静态 SID，不让 Qwen 每 batch 追逐变化标签。", end + 0.24)

    s = deck.slide("QG：Query 监督深度已经构建", "深度由 Train 的 Query→POI 类别分布决定，不读取旧 SID 定义，避免循环监督。",
                   "qg_prqk/README.md；IMPLEMENTATION_STATUS.md P2 / active P2.5-CAT")
    end = deck.table(s, ["深度", "可确定的信息", "Query 数", "允许监督的层"], [
        ["D0", "统计不足或目标分布不集中", "1,029,782", "不监督 SID；原 SFT 请求仍保留"],
        ["D1", "粗类别：19 类", "25,639", "只监督 S1"],
        ["D2", "细类别：全局402 / active397 类", "25,650", "监督 S1、S2"],
        ["D3", "高置信 Exact POI", "291,590", "监督 S1、S2、S3"],
    ], [1.25, 4.75, 2.01, 4.05])
    deck.takeaway(s, "共 1,372,661 个 normalized Query，展开 1,204,570 条层级边。D2 要求频次≥3、集中度≥0.80、归一熵≤0.40；D1 为≥3、≥0.85、≤0.35。", end + 0.25)
    deck.takeaway(s, "D3：Query / 正 pair 频次均≥2，Top-1 目标占比≥0.75、与 Top-2 差≥0.25、归一熵≤0.55。粗类 Query 只到 S1，细类只到 S2，不被强迫区分具体门店。", end + 1.10)

    s = deck.slide("QG：Exact Query Adapter 怎么训练", "只调整 D3 Query view，BGE 主干和 active POI 向量全部冻结。",
                   "qg_prqk/src/qg_prqk/query_adapter.py；METHOD_SPEC §5；P3A-FULL 已完成")
    end = deck.table(s, ["部分", "具体实现"], [
        ["残差结构", "BGE(q) → LayerNorm → Linear(1024,64) → GELU → Dropout → Linear(64,1024)"],
        ["输出", "与原始 Query 向量相加，再 L2 normalize；末层零初始化，从 identity 开始"],
        ["监督", "D3 Query 对应精确正 POI；对比学习，屏蔽重复正目标与已知假负例"],
        ["困难负例", "每条 Query 固定16条：语义6、词法4、局部地理6；另有动态 in-batch 负例"],
        ["用途边界", "只服务离线 SID 构建；不是 Qwen 在线前处理，不改写原始 SFT Query"],
    ], [2.10, 9.96])
    deck.takeaway(s, "SELECT：271,590 条训练 + 20,000 条 Train 内部 holdout 选 epoch。FINAL：从共同初始状态，用全部 291,590 条固定重训两轮，不续训 SELECT 或旧 Gate。", end + 0.25)

    s = deck.slide("QG：P3A-FULL 已完成，内部验证正向", "同一 active 候选库与 Train 内部 holdout，FULL-SELECT 的 Recall@10 为 77.100%。",
                   "query_adapter_exact_full/comparison/three_view_comparison.json；以下单位 %")
    rows = []
    for label, key in [("Raw BGE", "raw_bge_query"), ("历史 50k Gate", "gate50k_adapter"), ("FULL-SELECT epoch 2", "p3a_full_select_adapter")]:
        m = comparison["views"][key]
        assert m["rows"] == 20000
        rows.append([label] + [pct(m[k], 3) for k in ("recall_at_1", "recall_at_10", "mrr_at_10", "difficult_recall_at_10")])
    end = deck.table(s, ["Query view", "Recall@1", "Recall@10", "MRR@10", "困难 R@10"], rows, [3.46, 2.15, 2.15, 2.15, 2.15], highlight=2)
    deck.takeaway(s, "相对 Raw：R@1 +5.360pp、R@10 +3.115pp、困难 R@10 +6.147pp。困难集固定为 Raw Top-1 未命中的 10,281 条，三组共用假负例 mask。", end + 0.28)
    deck.takeaway(s, "这是 Adapter 的 Train 内部验证，不是业务固定 10k，不是 SID 或 SFT 提升。FINAL 包含原 holdout，不再在该集合宣称独立验证；当前仅冻结 FINAL 为 D3 默认 Query view。", end + 1.13)

    s = deck.slide("QG：双视图共享 token 的量化设计", "同一个 SID token 对应两个中心，用图约束对齐分配，不要求 Query 与 POI 共用向量中心。", src)
    end = deck.table(s, ["环节", "计划如何实现"], [
        ["输入视图", "POI 与 Query 应用同一 POI 拟合的公共方向去除；D1/D2 用 Raw，D3 用 FINAL Adapter"],
        ["双质心", "每层各有 1,024 个 POI 中心 U 与 Query 中心 V；U[k]、V[k] 共用 token k"],
        ["S1 目标", "内容距离 + D1/D2/D3 Query 距离 + Query–POI 分配不一致惩罚 + 粗类别软代价"],
        ["S2 目标", "残差距离 + D2/D3 对齐；用 P(细类 | S1,S2)，不直接硬设 S2=类别"],
        ["交替更新", "交替更新两侧分配与质心；稀疏 Query 中心向 POI 中心收缩，软权重逐步打开"],
    ], [2.10, 9.96])
    deck.takeaway(s, "下一层残差：r_next = normalize(r − Proj_center(r))，POI / Query 分别计算；不在每层重复使用原始 Query。P4–P6 尚未实现或正式运行。", end + 0.25)

    s = deck.slide("QG：S3 局部地理与困难实体设计", "地理只参与末层局部消歧；不把经纬度文本拼入 BGE，也不干扰 S1/S2 粗层组织。", src)
    end = deck.table(s, ["部分", "计划如何实现"], [
        ["局部父组", "按 (GID6,S1,S2) 分组；GID 是静态分组条件，不参与语义残差相减"],
        ["位置表示", "北京米制坐标：cell-relative dx/dy、父组距离 log(1+d)、方位 sin/cos"],
        ["稳定处理", "稳健标准化再 L2；单例父组的局部地理项与 gate 置零"],
        ["S3 目标", "POI / D3 Query 残差对齐 + 局部地理失真 + 困难实体同码惩罚"],
        ["困难边", "同父组、同细类、语义/名称近似；每 POI 最多20条；保护已知假负例"],
    ], [2.10, 9.96])
    deck.takeaway(s, "不为了静态唯一率强制父组内每个 POI 都使用不同 S3。这里的地理是离线 POI 侧设计，不是当前用户位置输入；P7 尚未实现和验证。", end + 0.25)

    s = deck.slide("QG：当前停在 Adapter 验收，尚无 SID", "P3A-FULL 已完成；方法线仍在推进，但当前不是后台训练运行中。",
                   "qg_prqk/docs/experiments/QG_PRQK_IMPLEMENTATION_STATUS.md；2026-09-06 终态")
    end = deck.table(s, ["阶段", "当前状态", "下一步产物"], [
        ["P2 / active P2.5", "已完成并验收", "Query 统计、类别深度与 active 行映射"],
        ["P3A-FULL", "SELECT / FINAL 完成，独立验证通过", "D3 默认 Query Adapter，72项测试记录"],
        ["P4 / P5", "待 P3A-FULL 审核；未启动", "层级图与 Query view；POI-only 内部 A0"],
        ["P6 / P7", "设计已记录；未实现 / 未运行", "类别+Query 的 S1/S2；Local Geo S3"],
        ["P8 与下游", "未运行；SID 静态验收后再审核", "SID / Bucket 产物；Final PID 与 SFT 延后"],
    ], [2.2, 4.75, 5.11])
    deck.takeaway(s, "QG 当前仍固定 1024×3，active 两组基线为 512×3。未来若归因比较，需要先对齐容量等协议；本次汇报不擅自改 QG 配置或启动下一阶段。", end + 0.25)


def render_preview(prs, output: Path) -> list[str]:
    """Render text, tables and native flow geometry; not an Office-rendered export."""
    output.mkdir(parents=True, exist_ok=True)
    issues: list[str] = []
    pages = []
    scale = 144

    def paint_text(draw, frame, box, label):
        x, y, width, height = box
        left = x + frame.margin_left / 914400 * scale
        top = y + frame.margin_top / 914400 * scale
        usable = width - (frame.margin_left + frame.margin_right) / 914400 * scale
        lines_to_draw = []
        for paragraph in frame.paragraphs:
            size = next((r.font.size.pt for r in paragraph.runs if r.font.size), 16)
            run = paragraph.runs[0] if paragraph.runs else None
            color = INK
            if run:
                try:
                    color = str(run.font.color.rgb)
                except (AttributeError, TypeError):
                    pass
            spacing = paragraph.line_spacing if isinstance(paragraph.line_spacing, float) else 1.2
            texts = wrapped(paragraph.text, usable / scale * 72, size)
            for line in texts:
                lines_to_draw.append((line, size, color, paragraph.alignment, spacing))
            if paragraph.space_after:
                lines_to_draw.append(("", paragraph.space_after.pt / 1.2, color, PP_ALIGN.LEFT, 1.2))
        needed = sum(size*spacing/72*scale for _,size,_,_,spacing in lines_to_draw)
        available = height - (frame.margin_top + frame.margin_bottom) / 914400 * scale
        if needed > available + 4:
            issues.append(f"{label}: 估计高度 {needed:.0f}px > {available:.0f}px")
        if frame.vertical_anchor == MSO_ANCHOR.MIDDLE:
            top += max(0, (available-needed)/2)
        for line, size, color, alignment, spacing in lines_to_draw:
            ft = font(size)
            xpos = left
            if alignment == PP_ALIGN.CENTER:
                xpos += (usable-ft.getlength(line))/2
            elif alignment == PP_ALIGN.RIGHT:
                xpos += usable-ft.getlength(line)
            draw.text((xpos, top), line, fill="#"+color, font=ft, anchor="lt")
            top += size*spacing/72*scale

    for i, slide in enumerate(prs.slides, 1):
        picture = Image.new("RGB", (1920, 1080), "white")
        draw = ImageDraw.Draw(picture)
        # The canonical master supplies a navy-to-blue title band. Pillow does
        # not render slide masters, so reproduce that band for geometry review.
        if i > 1:
            band_h = round(762000 / 914400 * scale)
            left = (26, 58, 107)
            right = (46, 124, 246)
            for y in range(band_h):
                ratio = y / max(1, band_h - 1)
                color = tuple(round(a + (b-a)*ratio) for a, b in zip(left, right))
                draw.line((0, y, 1920, y), fill=color)
            draw.rectangle((0, 0, round(57150/914400*scale), band_h), fill="#00C2A8")
            draw.rectangle((0, band_h, 1920, band_h+5), fill="#E5E7EB")
        for shape in slide.shapes:
            box = [v/914400*scale for v in (shape.left,shape.top,shape.width,shape.height)]
            if shape._element.tag.endswith("}cxnSp"):
                x1,y1,x2,y2 = [v/914400*scale for v in
                               (shape.begin_x, shape.begin_y, shape.end_x, shape.end_y)]
                draw.line((x1,y1,x2,y2), fill="#64748B", width=2)
                if shape._element.xpath(".//a:tailEnd[@type='triangle']"):
                    angle = math.atan2(y2-y1, x2-x1)
                    ux, uy = math.cos(angle), math.sin(angle)
                    draw.polygon([(x2,y2), (x2-10*ux+4*uy,y2-10*uy-4*ux),
                                  (x2-10*ux-4*uy,y2-10*uy+4*ux)], fill="#64748B")
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                x,y,w,h = box
                source = Image.open(BytesIO(shape.image.blob)).convert("RGBA")
                source = source.resize((max(1, round(w)), max(1, round(h))), Image.Resampling.LANCZOS)
                picture.paste(source, (round(x), round(y)), source)
            elif shape.has_text_frame:
                if shape.name in {"flow_node", "method_node", "method_token", "method_formula_bg"}:
                    x,y,w,h = box
                    try:
                        fill = str(shape.fill.fore_color.rgb)
                    except (AttributeError, TypeError):
                        fill = "FFFFFF"
                    try:
                        outline = str(shape.line.color.rgb)
                    except (AttributeError, TypeError):
                        outline = "CBD5E1"
                    draw.rectangle((x,y,x+w,y+h), fill="#"+fill, outline="#"+outline, width=1)
                paint_text(draw, shape.text_frame, box, f"第{i}页 {shape.name}")
            elif shape.has_table:
                y = box[1]
                for ri, row in enumerate(shape.table.rows):
                    h = row.height/914400*scale
                    x = box[0]
                    for ci, cell in enumerate(row.cells):
                        w = shape.table.columns[ci].width/914400*scale
                        try:
                            fill = str(cell.fill.fore_color.rgb)
                        except (AttributeError, TypeError):
                            fill = HEADER if ri == 0 else "FFFFFF"
                        draw.rectangle((x,y,x+w,y+h), fill="#"+fill, outline="#CBD5E1", width=1)
                        paint_text(draw, cell.text_frame, [x,y,w,h], f"第{i}页表{ri},{ci}")
                        x += w
                    y += h
        picture.save(output/f"slide-{i:02d}.png")
        pages.append(picture.resize((640,360)))
    for offset in range(0, len(pages), 9):
        batch = pages[offset:offset+9]
        sheet = Image.new("RGB", (1920, 3*395), "#E2E8F0")
        d = ImageDraw.Draw(sheet)
        for j, picture in enumerate(batch):
            x,y = (j%3)*640, (j//3)*395
            sheet.paste(picture,(x,y+30))
            d.text((x+10,y+2), str(offset+j+1), fill="#111827", font=font(10))
        sheet.save(output/f"contact-{offset//9+1:02d}.png")
    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description="复用0810模板，生成0907实验进展PPT并检查版面。")
    parser.add_argument("--output", type=Path, default=ROOT / "PPT/胡丹-0907.pptx")
    parser.add_argument("--preview-dir", type=Path, default=ROOT / "outputs/ppt_0907/preview")
    parser.add_argument("--reference-preview", action="store_true", help="仅生成参考PPT的文本表格预览")
    args = parser.parse_args()
    if not FONT_PATH.is_file():
        parser.error(f"缺少预览中文字体：{FONT_PATH}")
    if args.reference_preview:
        for name in ("0727", "0810", "0824"):
            render_preview(Presentation(ROOT/f"PPT/胡丹-{name}.pptx"), args.preview_dir/name)
        return 0
    deck = Deck()
    from method_story_0907 import build_method_story
    build_method_story(deck)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    deck.prs.save(args.output)
    with ZipFile(args.output) as z:
        assert z.testzip() is None
    reopened = Presentation(args.output)
    issues = render_preview(reopened, args.preview_dir)
    print(json.dumps({"output":str(args.output), "slides":len(reopened.slides),
                      "preview":"Pillow geometry preview, not Office rendering", "layout_warnings":issues},
                     ensure_ascii=False, indent=2))
    return 1 if issues else 0


if __name__ == "__main__":
    raise SystemExit(main())
