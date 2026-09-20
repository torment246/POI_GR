"""Editable method flows for the 0907 deck; result tables stay unchanged."""

from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


FLOW_SLIDES = (2, 4, 6, 8, 10, 11, 15, 17, 18, 19, 20, 22, 23, 24, 25)


class Flow:
    """Small layout helpers using the deck's original typography and margins."""

    def __init__(self, deck, page: int):
        self.deck = deck
        self.slide = deck.prs.slides[page - 1]
        for shape in list(self.slide.shapes):
            # Keep title, page number, claim and source; replace only the body.
            if Inches(1.8) <= shape.top < Inches(6.90):
                shape._element.getparent().remove(shape._element)

    def text(self, text: str, x: float, y: float, w: float, h: float = 0.65,
             *, size: float = 16, bold: bool = False, muted: bool = False,
             center: bool = False):
        return self.deck.text(self.slide, text, y, h, x=x, w=w, size=size,
                              color="64748B" if muted else "111827", bold=bold,
                              align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT)

    def node(self, text: str, x: float, y: float, w: float, h: float = 0.90,
             *, emphasis: bool = False):
        shape = self.text(text, x, y, w, h, center=True)
        shape.name = "flow_node"
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string("EAF2FF" if emphasis else "F8FAFC")
        shape.line.color.rgb = RGBColor.from_string("CBD5E1")
        shape.line.width = Pt(0.75)
        frame = shape.text_frame
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = frame.margin_right = Inches(0.10)
        frame.margin_top = frame.margin_bottom = Inches(0.07)
        for p in frame.paragraphs:
            p.space_after = Pt(0)
        return shape

    def path(self, points: list[tuple[float, float]], *, arrow: bool = True):
        for index, (start, end) in enumerate(zip(points, points[1:])):
            shape = self.slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, Inches(start[0]), Inches(start[1]),
                Inches(end[0]), Inches(end[1]))
            shape.name = "flow_edge"
            shape.line.color.rgb = RGBColor.from_string("64748B")
            shape.line.width = Pt(1.15)
            if arrow and index == len(points) - 2:
                tail = OxmlElement("a:tailEnd")
                tail.set("type", "triangle")
                tail.set("w", "med")
                tail.set("len", "med")
                shape.line._get_or_add_ln().append(tail)

    def chain(self, labels: list[str], y: float, *, x: float = 0.65,
              w: float = 12.0, h: float = 0.90):
        gap = 0.50
        width = (w - gap * (len(labels) - 1)) / len(labels)
        for i, label in enumerate(labels):
            pos = x + i * (width + gap)
            self.node(label, pos, y, width, h)
            if i:
                self.path([(pos - gap + 0.03, y + h / 2), (pos - 0.04, y + h / 2)])

    def note(self, text: str, y: float = 6.03, h: float = 0.74):
        self.text(text, 0.65, y, 12.0, h)


def overview(f: Flow) -> None:
    f.chain(["已完成：效果诊断", "当前：active 配对基线", "新方法：QG-PRQK"], 2.15)
    f.text("桶内重排与 TIGER-Joint\nGID 顺序与 BeamRisk-SFT\n正式评测均未超过 TIGER", 0.65, 3.35, 3.67, 1.35)
    f.text("716,245 条活跃 POI\nBGE / MMBERT，512×3\nSID 与 SFT 数据完成\n两组均等待平台训练", 4.82, 3.35, 3.67, 1.65)
    f.text("分层 Query 监督\n双视图量化与局部地理\nExact Adapter 已完成\n后续 SID 待评审、未启动", 8.99, 3.35, 3.67, 1.65)
    f.note("判断标准保持不变：最终生成式 HR / NDCG。静态 SID 质量、Adapter 向量召回、合法路径比例，分别报告，不替代 SFT 效果。")


def bucket(f: Flow) -> None:
    f.chain(["冻结 epoch 3 Qwen\n无约束 Beam=10", "抽取 S1 / S2 / S3\n合法桶按首次出现去重", "冻结 mapping 展开\n桶内全部 POI"], 2.02)
    f.path([(10.82, 2.92), (10.82, 3.37), (6.65, 3.37), (6.65, 3.68)])
    f.node("热度：Train log1p(count)\n词法：Query 对名称 / 地址\n语义：冻结向量点积", 4.68, 3.71, 3.94, 1.23)
    f.node("排序后精确 POI Top-10\n对比原完整 SID+C", 9.20, 3.71, 3.43, 1.23)
    f.path([(8.62, 4.32), (9.17, 4.32)])
    f.text("排序对照\n保持桶序、仅桶内排序\n全局排序 / RRF 融合\n共八种变体", 0.65, 3.68, 3.60, 1.60)
    f.note("不改变 SID、Qwen 或原始 Beam，不新增 SFT。E4 组仅替换 POI 语义向量；两组均不用地理特征、不训练排序器、不在 Validation 调参。")


def joint(f: Flow) -> None:
    f.node("冻结 BGE POI 向量\n目标 / 历史 POI 行号", 0.65, 2.10, 3.22)
    f.node("当前 RQ-VAE\nfresh KMeans，1024×3", 4.45, 2.10, 3.45)
    f.node("实时 hard SID 标签\nargmin 阻断生成梯度", 9.05, 2.10, 3.58)
    f.path([(3.87, 2.55), (4.42, 2.55)])
    f.path([(7.90, 2.55), (9.02, 2.55)])
    f.node("Qwen 基座重新扩词\n请求 hidden 投影到 256D", 0.65, 4.03, 3.70, 1.05)
    f.node("L_align：InfoNCE\n连接 Qwen / 投影 / encoder", 4.92, 4.03, 3.79, 1.05)
    f.node("L_gen：目标 SID CE\n只更新 Qwen", 9.28, 4.03, 3.35, 1.05)
    f.path([(4.35, 4.55), (4.89, 4.55)])
    f.path([(6.18, 3.00), (6.18, 4.00)])
    f.path([(10.84, 3.00), (10.84, 4.00)])
    f.path([(2.50, 4.03), (2.50, 3.57), (10.18, 3.57), (10.18, 4.00)])
    f.text("L_rq：独立均匀目录批做重建，更新 encoder / decoder / codebook。", 0.65, 5.32, 12.0, 0.49)
    f.note("总目标 L = L_gen + 0.1 L_align + L_rq。SID 随更新变化；epoch 3 冻结导出目录，只评三级 Bucket，不追加 C。")


def gid_order(f: Flow) -> None:
    f.text("共同输入：当前 Query + 用户 GID + 最近 10 条历史；同一 POI / SID3 / GID6 / 局部 D", 0.65, 1.95, 12.0, 0.61)
    f.text("GID-first", 0.65, 2.86, 1.65, 0.42, bold=True)
    f.chain(["G1…G6", "S1 → S2 → S3", "[D] 去重"], 2.72, x=2.48, w=10.15)
    f.text("先选区域，再在区域内找语义实体", 2.48, 3.75, 10.15, 0.44)
    f.text("SID-first", 0.65, 4.51, 1.65, 0.42, bold=True)
    f.chain(["S1 → S2 → S3", "G1…G6", "[D] 去重"], 4.37, x=2.48, w=10.15)
    f.text("先选语义，再用地理定位实体。两版目标均为 9/10 个编码 token。", 2.48, 5.40, 10.15, 0.46)
    f.note("历史与目标一起改序；扩词初始化、样本和三轮 SFT 相同。不同 SID3 / POI 数为 71.68%，加 GID6 为 82.76%，局部 D 后 100% 唯一。")


def beam_risk(f: Flow) -> None:
    f.node("当前模型 Beam=10\n追踪 gold 路径位置", 0.65, 3.22, 3.13, 1.03)
    cases = [
        (1.99, "首次掉出 Beam", "Survive：gold 前缀\n对比当层实际第 10 名边界前缀"),
        (3.44, "进入第 2–10 名", "Rank：完整 gold SID\n对比当前第 1 名错误 SID"),
        (4.89, "已经排第 1 名", "不额外构造风险 pair\n只保留标准 CE"),
    ]
    for y, label, target in cases:
        f.node(label, 4.53, y, 2.68, 0.90)
        f.node(target, 8.0, y, 4.63, 0.90)
        f.path([(3.78, 3.74), (4.13, 3.74), (4.13, y + 0.45), (4.50, y + 0.45)])
        f.path([(7.21, y + 0.45), (7.97, y + 0.45)])
    f.note("每请求仅取一种风险。Score 为 token log 概率之和；L_risk = softplus(Score_neg − Score_gold)。该损失越过边界仍有梯度，是本轮暴露的问题。", 6.03)


def beam_training(f: Flow) -> None:
    f.text("主训练流：每轮完整遍历 7,586,410 条 Train，2,851,853 packed rows", 0.65, 1.95, 12.0, 0.52)
    f.chain(["Epoch 1\n普通 CE", "Epoch 2\n完整 CE + 风险梯度", "Epoch 3\n完整 CE + 刷新风险"], 2.70)
    f.node("自挖掘 1\nepoch-1 Beam 选 100k pair", 2.28, 4.25, 3.87, 0.94)
    f.node("自挖掘 2\n刷新下一轮的风险 pair", 7.62, 4.25, 3.87, 0.94)
    f.path([(2.48, 3.60), (2.48, 4.22)])
    f.path([(6.15, 4.72), (6.65, 4.72), (6.65, 3.63)])
    f.path([(7.10, 3.60), (7.10, 4.72), (7.59, 4.72)])
    f.path([(10.81, 4.25), (10.81, 3.63)])
    f.text("辅助流：Train-only 500k 候选池；每 4 optimizer step 注入一次，weight=0.2，全局 128 pair。", 0.65, 5.40, 12.0, 0.63)
    f.note("从扩词基座重训，恢复同一 optimizer / scheduler / 四 rank RNG；三轮共 16,713 step。只评 epoch 3 固定 10k：无约束 + 合法路径约束，不追加第四轮。", 6.11, 0.64)


def active_sid(f: Flow) -> None:
    f.text("共同输入：716,245 条 active POI；同一行序、同一名称 / 地址 / 别名文本", 0.65, 1.96, 12.0, 0.49)
    f.chain(["冻结 BGE-M3\n1,024 维", "RQ-VAE\n1024 → 512 → 256", "SID：S1 / S2 / S3\n512×512×512", "桶内稳定编号 C\n唯一四层标识"], 2.69)
    f.chain(["冻结 MMBERT\n池化 + 768→128 + L2", "RQ-VAE\n128 → 512 → 256", "SID：S1 / S2 / S3\n512×512×512", "桶内稳定编号 C\n唯一四层标识"], 4.02)
    f.text("两路均训练 20 epoch，seed=42；全部 716,245 向量参与初始化；C 按 POI ID 稳定分配。", 0.65, 5.15, 12.0, 0.64)
    f.note("活跃闭集由两周全部目标与历史身份定义；量化不使用评测 Query–POI 监督。缩库改变候选协议，不能当作未来开放目录评测。")


def active_sft(f: Flow) -> None:
    f.chain(["样本只替换 identifier\nTrain / Valid / Test 不变", "安全检查与 packed 缓存\n两组输入均已就绪", "4×6000D，三轮全参\nQwen3-0.6B"], 2.05)
    f.text("原始样本 7,586,410 / 597,421 / 606,682；packed Train / Valid = 1,296,883 / 96,949。\ncutoff=1024，最长 953，超长 / 目标截断均为 0；batch 8 × 累积16 × 4卡 = 512。", 0.65, 3.15, 12.0, 0.89)
    f.node("训练结束\n只取 epoch 3", 10.13, 4.14, 2.50, 0.94)
    f.node("固定 Validation 10k\n无约束 + 全目录合法路径约束", 0.65, 4.14, 8.25, 0.87)
    f.node("四类泛化集，各 10k，均无约束\n新 Pair / 新 Query / 长尾 / 冷目标", 0.65, 5.22, 8.25, 0.87)
    f.path([(11.82, 2.95), (12.85, 2.95), (12.85, 4.61), (12.66, 4.61)])
    for y in (4.575, 5.655):
        f.path([(10.13, 4.61), (9.54, 4.61), (9.54, y), (8.93, y)])
    f.note("已通过 dry-run，尚未平台实训。新 Pair 为已见 Query / 新 Pair；新 Query 为新 Query / 已见 POI。两组同目录配对比较，不混入旧全库结论。", 6.22, 0.61)


def qg_overview(f: Flow) -> None:
    f.chain(["Train Query 统计\nD1 / D2 / D3 分层", "POI / Query 双视图\n独立向量与质心", "S1 / S2 粗到细\n类别软锚点与图对齐", "S3 局部消歧\n相对地理与困难实体"], 2.23)
    f.text("已完成\nQuery 深度与 active 映射\nD3 Exact Adapter", 0.65, 3.52, 2.69, 1.21)
    f.text("计划设计\n同一个 token 对应双中心\n在离散分配中对齐", 3.82, 3.52, 2.70, 1.21)
    f.text("计划设计\n只监督 Query 能确定的层\n不硬绑类别 ID", 6.99, 3.52, 2.69, 1.21)
    f.text("计划设计\n仅在末层引入地理\n不强制 S3 唯一", 10.16, 3.52, 2.47, 1.21)
    f.note("与 E4 不同：不把全部 Query 混入单一 POI 向量。与 Joint 不同：先离线构建静态 SID，不让 Qwen 每个 batch 追逐变化标签。P4–P8 尚未启动。")


def qg_depth(f: Flow) -> None:
    f.node("Train Query → POI\n类别分布与精确目标置信度", 0.65, 3.02, 3.14, 1.08)
    branches = [
        (1.95, "D0：不足 / 不集中", "1,029,782 个；不监督 SID"),
        (2.87, "D1：粗类别（19 类）", "25,639 个；只到 S1"),
        (3.79, "D2：细类别（402 / 397）", "25,650 个；到 S1 + S2"),
        (4.71, "D3：高置信 Exact POI", "291,590 个；到 S1 + S2 + S3"),
    ]
    for y, label, scope in branches:
        f.node(label, 4.55, y, 3.91, 0.68)
        f.text(scope, 9.05, y + 0.05, 3.60, 0.61)
        f.path([(3.79, 3.56), (4.15, 3.56), (4.15, y + 0.34), (4.52, y + 0.34)])
        f.path([(8.46, y + 0.34), (9.02, y + 0.34)])
    f.text("共 1,372,661 个 normalized Query、1,204,570 条层级边；D0 原 SFT 请求仍保留。\nD1：频次≥3、集中度≥0.85、归一熵≤0.35；D2：≥3、≥0.80、≤0.40。", 0.65, 5.41, 12.0, 0.72)
    f.note("D3：Query / 正 pair 频次均≥2，Top-1 占比≥0.75，与 Top-2 差≥0.25，归一熵≤0.55。\nD2 类别数为全局 402 / active 397。不读取旧 SID 定义；粗类不被强迫区分具体门店。", 6.16, 0.69)


def qg_adapter(f: Flow) -> None:
    f.chain(["冻结 BGE(q)\n1,024D Query", "LN → Linear(64)\nGELU → Dropout", "Linear(1024)\n末层零初始化", "加原 Query，再 L2\nD3 Query view"], 2.07)
    f.path([(1.96, 2.97), (1.96, 3.33), (11.32, 3.33), (11.32, 3.00)])
    f.text("残差支路：从 identity 开始；BGE 主干与 active POI 向量均冻结", 3.18, 3.47, 8.80, 0.47)
    f.node("D3 精确正 POI + 固定16负例\n语义6 / 词法4 / 局部地理6", 0.65, 4.22, 4.98, 0.98)
    f.node("对比学习\n另加 in-batch 负例，屏蔽已知假负例", 6.21, 4.22, 6.42, 0.98)
    f.path([(5.63, 4.71), (6.18, 4.71)])
    f.path([(11.75, 2.97), (12.86, 2.97), (12.86, 4.71), (12.66, 4.71)])
    f.text("SELECT：271,590 训练 + 20,000 Train 内部 holdout，选 epoch 2。\nFINAL：从共同初始状态，用全部 291,590 条固定重训两轮，不续训 SELECT / 旧 Gate。", 0.65, 5.42, 12.0, 0.73)
    f.note("仅服务离线 SID 构建，不是 Qwen 在线前处理，不改原始 SFT Query；P3A-FULL 已完成。", 6.32, 0.46)


def qg_quantizer(f: Flow) -> None:
    f.node("POI 残差视图", 0.65, 2.31, 2.43, 0.73)
    f.node("POI 质心 U[k]\n每层 1,024 个", 3.62, 2.18, 2.70, 0.99)
    f.node("Query 残差视图\nD1/D2 Raw，D3 Adapter", 0.65, 4.10, 2.95, 0.99)
    f.node("Query 质心 V[k]\n每层 1,024 个", 4.10, 4.10, 2.75, 0.99)
    f.node("共享离散 token k\n两侧分配一致性图约束", 7.53, 3.01, 5.10, 0.99)
    f.path([(3.08, 2.68), (3.59, 2.68)])
    f.path([(3.60, 4.60), (4.07, 4.60)])
    f.path([(6.32, 2.68), (6.97, 2.68), (6.97, 3.51), (7.50, 3.51)])
    f.path([(6.85, 4.60), (7.20, 4.60), (7.20, 3.75), (7.50, 3.75)])
    f.text("S1：内容 / Query 距离 + 粗类软代价\nS2：残差对齐 + P(细类 | S1,S2)\n不设 S2=类别，不强制 U[k]=V[k]", 7.53, 4.28, 5.10, 1.17)
    f.text("交替优化：两侧分配与中心", 7.53, 2.18, 5.10, 0.47)
    f.text("输入共用 POI 拟合的公共方向去除；稀疏 Query 中心向 POI 中心收缩，软权重逐步打开。", 0.65, 5.55, 12.0, 0.60)
    f.note("各视图分别递推：r_next = normalize(r − Proj_center(r))；不逐层重用原始 Query。\nS1 用 D1/D2/D3，S2 用 D2/D3。P4–P6 尚未实现或正式运行。", 6.16, 0.69)


def qg_geo(f: Flow) -> None:
    f.node("固定局部父组\n(GID6, S1, S2)", 0.65, 3.14, 2.45, 1.00)
    f.node("北京米制局部地理\ndx/dy、log(1+d)、方位 sin/cos", 3.88, 1.97, 4.77, 0.96)
    f.node("同父组、同细类困难实体\n语义 / 名称近似，最多20边 / POI", 3.88, 3.54, 4.77, 0.96)
    f.node("S3 联合分配\n残差对齐 + 地理失真\n+ 困难实体同码惩罚", 9.46, 3.02, 3.17, 1.26)
    for y in (2.45, 4.02):
        f.path([(3.10, 3.64), (3.49, 3.64), (3.49, y), (3.85, y)])
        f.path([(8.65, y), (9.04, y), (9.04, 3.65), (9.43, 3.65)])
    f.node("POI / D3 Query 的末层残差", 3.88, 5.08, 4.77, 0.70)
    f.path([(8.65, 5.43), (11.04, 5.43), (11.04, 4.31)])
    f.text("稳健标准化 + L2；单例父组地理项与 gate 置零；保护已知假负例，不强制组内 S3 唯一。\n地理仅用于离线 POI 侧，不参与语义残差相减；不是用户当前位置输入。P7 尚未实现。", 0.65, 5.90, 12.0, 0.94)


def qg_status(f: Flow) -> None:
    f.chain(["P2 / active P2.5\n已完成并验收", "P3A-FULL\nSELECT / FINAL 已完成", "审核关口\n当前停在这里"], 2.04)
    f.text("Query 统计、类别深度\nactive 行映射", 0.65, 3.10, 3.65, 0.83)
    f.text("D3 默认 Adapter\n独立检查通过，72 项测试记录", 4.82, 3.10, 3.65, 0.83)
    f.text("不是后台训练运行中\n审核后再推进以下阶段", 8.99, 3.10, 3.65, 0.83)
    f.chain(["P4 / P5：未启动\n层级图 / Query view / A0", "P6 / P7：未实现\n类别 + Query / Local Geo", "P8 及下游：未运行\nSID / Bucket → PID / SFT"], 4.42)
    f.path([(10.82, 2.94), (12.87, 2.94), (12.87, 4.14), (2.48, 4.14), (2.48, 4.39)])
    f.note("QG 当前设计仍为 1024×3，active 两组基线为 512×3。未来归因比较需先对齐容量等协议；本次不改配置、不启动后续阶段。")


def conclusion(f: Flow) -> None:
    f.text("已知结论", 0.65, 1.97, 2.0, 0.48, size=20, bold=True)
    f.text("桶内重排、Joint、GID 顺序与 BeamRisk v1 均未超过 TIGER；保留产物与基础设施。\n更低 loss、更高合法率或静态唯一率，都不能直接判定生成召回提升。", 0.65, 2.61, 12.0, 0.80)
    f.chain(["近期：active 配对结果\n同目录 / 512×3 / 同 SFT", "随后：QG Adapter 审核\n通过后推进离线 SID", "最终：生成效果归因\n对齐容量与评测协议"], 3.76)
    f.text("当前位置输入", 0.65, 5.03, 2.8, 0.48, size=20, bold=True)
    f.text("用户 GID 是输入条件，与输出 SID 独立；不要求给目标 SID 再拼 GID。", 0.65, 5.62, 12.0, 0.51)
    f.note("active 等待训练和固定 10k + 四类泛化评测；QG 只有 Adapter 内部验证正向，尚无 SID / SFT。新旧候选库、不同输出结构、不同评测集合分开报告。", 6.22, 0.62)


def apply_flow_revision(deck) -> None:
    """Replace method bodies only and normalize all major titles to true black."""
    layouts = dict(zip(FLOW_SLIDES, (overview, bucket, joint, gid_order, beam_risk,
        beam_training, active_sid, active_sft, qg_overview, qg_depth, qg_adapter,
        qg_quantizer, qg_geo, qg_status, conclusion)))
    for page, layout in layouts.items():
        layout(Flow(deck, page))
    for slide in deck.prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
            if any(r.font.size is not None and r.font.size.pt >= 20 for r in runs):
                for run in runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
