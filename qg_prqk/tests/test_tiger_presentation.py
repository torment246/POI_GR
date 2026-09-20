"""Synthetic checks for the frozen-catalog presentation comparison."""

import numpy as np
import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from qg_prqk.sid.tiger_presentation import collision_pairs, full_codes, replace_text
from qg_prqk.sid.tiger_local_cases import eligible_groups, gid6_string


def test_complete_sid_not_s3_alone() -> None:
    sid = np.array([[0, 0, 1], [1, 0, 1], [0, 0, 1], [0, 0, 1]])
    groups = np.array([0, 0, 0, 1])
    assert collision_pairs(groups, sid) == 1
    assert full_codes(sid)[0] != full_codes(sid)[1]


@pytest.mark.parametrize('sid', [np.zeros((2, 2), dtype=int),
                               np.array([[512, 0, 0]]),
                               np.array([[0, -1, 0]]),
                               np.array([[0., 1., 2.]])])
def test_invalid_sid_rejected(sid: np.ndarray) -> None:
    with pytest.raises(ValueError):
        full_codes(sid)


def test_replacement_preserves_editable_style() -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = '原文'
    run.font.name = 'Microsoft YaHei'
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor.from_string('FFFFFF')
    replace_text(box, 'TIGER\nQG-HRQ')
    assert box.text == 'TIGER\nQG-HRQ'
    for paragraph in box.text_frame.paragraphs:
        new = paragraph.runs[0]
        assert new.font.name == 'Microsoft YaHei'
        assert new.font.size == Pt(20)
        assert str(new.font.color.rgb) == 'FFFFFF'


def test_local_case_keeps_complete_group_and_one_prefix_per_method() -> None:
    groups = np.zeros(5, dtype=np.int64)
    tiger = np.tile([1, 2, 3], (5, 1))
    qg = np.column_stack([np.full(5, 4), np.full(5, 5), np.arange(5)])
    assert eligible_groups(groups, tiger, qg).tolist() == [0]
    tiger[0, 0] = 0
    assert eligible_groups(groups, tiger, qg).tolist() == []


def test_local_case_does_not_drop_colliding_or_excess_members() -> None:
    groups = np.zeros(9, dtype=np.int64)
    tiger = np.tile([1, 2, 3], (9, 1))
    qg = np.column_stack([np.full(9, 4), np.full(9, 5), np.arange(9)])
    assert eligible_groups(groups, tiger, qg).tolist() == []
    qg[0] = qg[1]
    assert eligible_groups(groups[:5], tiger[:5], qg[:5]).tolist() == []


def test_gid_is_packed_tokens_not_utf8() -> None:
    assert gid6_string(bytes([28, 29, 4, 12, 21, 18])) == 'wx4dpk'
    with pytest.raises(ValueError):
        gid6_string(b'wx4dpk')
