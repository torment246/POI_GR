"""Replace only the two local case slides with explicitly selected improvements."""

from __future__ import annotations

import argparse
from copy import deepcopy
import json
from pathlib import Path

import numpy as np
import pyarrow.parquet as pq
from pptx import Presentation
from pptx.util import Inches, Pt

from qg_prqk.artifacts import sha256_file, write_json_atomic
from qg_prqk.sid.local_visualization import local_coordinates
from qg_prqk.sid.geo import encode_geohash_tokens, geohash_strings
from qg_prqk.sid.tiger_presentation import (
    AFTER, NAMES, OUTPUT, ROOT, TIGER, WORK, checked, collision_pairs,
    draw_local_cases, frozen, full_codes, plot_setup, replace_picture,
    replace_text, validate_deck,
)

CASE_GROUPS = (55280, 221749)
CASE_TITLES = ('枣园小区：8 栋楼不再共享同一 SID', '菜市口地铁站：站点与出入口分开编码')


def gid6_string(value: bytes) -> str:
    """Decode the six packed integer tokens, not UTF-8 text."""
    tokens = np.frombuffer(value, dtype=np.uint8)
    if len(tokens) != 6 or np.any(tokens >= 32):
        raise ValueError('GID6 须为六个 0..31 的 token')
    return next(geohash_strings(tokens[None, :]))


def eligible_groups(groups: np.ndarray, tiger: np.ndarray, qg: np.ndarray) -> np.ndarray:
    """Find complete 5–8 POI groups with one prefix per method and clear improvement."""
    sizes = np.bincount(groups)
    pairs = []
    for sid in (tiger, qg):
        keys, counts = np.unique(groups.astype(np.int64) * 512**3 + full_codes(sid), return_counts=True)
        pairs.append(np.bincount(keys // 512**3, weights=counts*(counts-1)//2, minlength=len(sizes)))
    same_prefix = np.ones(len(sizes), dtype=bool)
    for sid in (tiger, qg):
        prefix = sid[:, 0].astype(np.int64)*512 + sid[:, 1]
        low, high = np.full(len(sizes), 512**2), np.zeros(len(sizes), dtype=np.int64)
        np.minimum.at(low, groups, prefix)
        np.maximum.at(high, groups, prefix)
        same_prefix &= low == high
    return np.flatnonzero((sizes >= 5) & (sizes <= 8) & (pairs[0] >= 6) & (pairs[1] == 0) & same_prefix)


def analyze_cases() -> dict:
    """Bind selected full-parent cohorts to the previously verified common catalog."""
    sources = {}
    prior_path = checked(WORK/'analysis.json', sources)
    prior = json.loads(prior_path.read_text())
    sid_manifest = TIGER/'sid_manifest.json'
    checked(sid_manifest, sources, prior['sources'][str(sid_manifest.relative_to(ROOT))])
    tiger_path = TIGER/'sid_codes.npy'
    tiger = np.load(checked(tiger_path, sources, prior['sources'][str(tiger_path.relative_to(ROOT))]))
    qg_path = frozen(AFTER, 'poi_sid_s1_s2_s3.npy', sources)
    meta_path = frozen(AFTER, 'poi_geo_metadata.parquet', sources)
    for path in (qg_path, meta_path):
        if sources[str(path.relative_to(ROOT))] != prior['sources'][str(path.relative_to(ROOT))]:
            raise ValueError('POI 对齐来源与此前验证不一致')
    if not prior['id_order_verified']:
        raise ValueError('缺少 POI ID 行序对齐验证')
    qg = np.load(qg_path)
    groups = np.load(frozen(AFTER, 'poi_parent_groups.npy', sources))
    if tiger.shape != qg.shape or qg.shape != (716245, 3) or groups.shape != (716245,):
        raise ValueError('目录规模不一致')
    eligible = eligible_groups(groups, tiger, qg)
    if not set(CASE_GROUPS).issubset(set(eligible.tolist())):
        raise ValueError('选例不满足公开的选择条件')
    cases = []
    for group, title in zip(CASE_GROUPS, CASE_TITLES):
        members = pq.read_table(meta_path, columns=[
            'poi_row_index', 'poi_id', 'displayname', 'category', 'lat', 'lng', 'gid6', 'parent_group'],
            filters=[('parent_group', '=', group)]).to_pylist()
        members.sort(key=lambda member: member['poi_row_index'])
        rows = np.asarray([m['poi_row_index'] for m in members])
        if not np.array_equal(rows, np.flatnonzero(groups == group)):
            raise ValueError('禁止丢弃局部父组成员')
        gids = {gid6_string(m['gid6']) for m in members}
        coordinate_gids = set(geohash_strings(encode_geohash_tokens(
            np.asarray([m['lng'] for m in members]), np.asarray([m['lat'] for m in members]))))
        if gids != coordinate_gids:
            raise ValueError('GID6 与真实经纬度不一致')
        if len(gids) != 1:
            raise ValueError('案例必须位于同一 GID6')
        xy = local_coordinates(np.asarray([m['lng'] for m in members]), np.asarray([m['lat'] for m in members]))
        distances = np.linalg.norm(xy[:, None]-xy[None], axis=2)
        case = dict(title=title, parent_group=group, gid6=gids.pop(), members=members,
                    max_distance_m=float(distances.max()),
                    min_distance_m=float(distances[np.triu_indices(len(rows), 1)].min()))
        for name, sid in zip(NAMES, (tiger, qg)):
            case[name] = dict(collision_pairs=collision_pairs(np.zeros(len(rows), dtype=int), sid[rows]),
                              prefix=sid[rows[0], :2].tolist(), unique_sids=len(np.unique(full_codes(sid[rows]))))
            for member, row in zip(members, rows):
                member[name] = sid[row].tolist()
                member['gid6'] = case['gid6']
        cases.append(case)
    return dict(sources=sources, geographic_cases=cases, selection=dict(
        eligible_count=len(eligible), eligible_groups=eligible.tolist(), chosen_groups=list(CASE_GROUPS),
        uses_outcome=True,
        criteria='Whole QG GID6/S1/S2 parent, 5..8 POIs, one S1/S2 per method, TIGER >=6 colliding pairs, QG=0.',
        choice='Manually selected for clear building/entrance names; first case is maximum collision reduction among eligible groups.'),
        limits=['Outcome-selected illustrations, not average effect or a Geo-only ablation.',
                'SID means three semantic tokens before C/GID/Dedup; no inference or training.',
                'Whole selected QG parents, not all POIs in the neighborhood or full TIGER collision buckets.'])


def update_cases(source: Path, output: Path, work: Path, data: dict) -> None:
    """Keep the existing two-column case layout and replace slides 17–18 only."""
    source_hash = sha256_file(source)
    deck = Presentation(source)
    if len(deck.slides) != 26:
        raise ValueError('输入必须是 26 页 TIGER 对比版')
    for i, case in enumerate(data['geographic_cases']):
        slide = deck.slides[16+i]
        replace_text(slide.shapes[0], case['title'])
        replace_text(slide.shapes[1], f"GID6：{case['gid6']}；改善选例，非平均效果")
        replace_text(slide.shapes[2], '组内 S1/S2 相同；同色 / 连线 = 同码')
        replace_picture(slide, 3, work/f'local_case_{i+1}.png')
        for col, name in enumerate(NAMES):
            values = case[name]
            prefix = '/'.join(map(str, values['prefix']))
            replace_text(slide.shapes[4+col],
                         f"S1/S2：{prefix}；{values['unique_sids']} 个 SID；同码对 {values['collision_pairs']}")
        label_template = deepcopy(slide.shapes[6]._element)
        for shape in list(slide.shapes)[6:]:
            slide.shapes._spTree.remove(shape._element)
        for j, member in enumerate(case['members']):
            element = deepcopy(label_template)
            slide.shapes._spTree.insert_element_before(element, 'p:extLst')
            label = slide.shapes[-1]
            element.xpath('.//p:cNvPr')[0].set('id', str(max(s.shape_id for s in slide.shapes)+1))
            col, row = divmod(j, 4)
            label.left, label.top = Inches(.64+col*6.16), Inches(6.31+row*.27)
            label.width, label.height = Inches(6.0), Inches(.28)
            text = f"{j+1}  {member['displayname']}   S3：{member['TIGER'][2]} → {member['QG-HRQ'][2]}"
            replace_text(label, text)
            for paragraph in label.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(15)
    deck.save(output)
    validate_deck(source, output, work, source_hash, {16, 17})


def main() -> int:
    parser = argparse.ArgumentParser(description='筛选明显的局部实体区分案例，只替换 TIGER 对比 PPT 第 17–18 页。')
    parser.add_argument('--source', type=Path, default=OUTPUT, help='26 页 TIGER 对比版 PPT')
    parser.add_argument('--output', type=Path, default=ROOT/'PPT/胡丹-0917-QG-HRQ-TIGER局部案例增强版.pptx', help='新 PPT，不覆盖')
    parser.add_argument('--work-dir', type=Path, default=ROOT/'qg_prqk/outputs/figures/tiger_local_cases_v2', help='新的分析和图目录')
    args = parser.parse_args()
    try:
        if args.output.exists() or args.work_dir.exists():
            raise ValueError('输出已存在，请使用新的文件和目录')
        if not args.work_dir.resolve().is_relative_to(ROOT/'qg_prqk/outputs'):
            raise ValueError('分析必须保存在 qg_prqk/outputs 内')
        data = analyze_cases()
        args.work_dir.mkdir(parents=True)
        plot_setup()
        draw_local_cases(data, args.work_dir)
        data['artifacts'] = {p.name: sha256_file(p) for p in args.work_dir.glob('*.png')}
        data['implementation_sha256'] = sha256_file(Path(__file__))
        write_json_atomic(args.work_dir/'analysis.json', data)
        update_cases(args.source, args.output, args.work_dir, data)
        for case in data['geographic_cases']:
            print(case['title'], {name: case[name] for name in NAMES}, flush=True)
        print(args.output)
    except (OSError, ValueError, KeyError) as error:
        parser.exit(2, f'局部案例更新失败：{error}\n')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
