"""Recompute the 0917 visual comparison against the frozen active-catalog TIGER."""

from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import pyarrow.parquet as pq
from lxml import etree

from qg_prqk.artifacts import sha256_file, write_json_atomic
from qg_prqk.sid.category_region_visualization import _weighted_purity
from qg_prqk.sid.local_visualization import AFTER, ROOT, local_coordinates, same_code_pairs
from qg_prqk.sid.query_visualization import GRAPH, METRICS, distribution, prefix_codes, select_case

TIGER = ROOT / 'outputs/sid/tiger/active_716k_bge_m3/TIGER-ACTIVE716K-BGE-M3-512x3-FULLINIT/evaluations/epoch_20'
EMBEDDING = ROOT / 'outputs/embeddings/beijing_poi_active_order14d_history10_bge_m3'
OLD_QUERY = ROOT / 'qg_prqk/outputs/figures/qg_prqk_query_positive_tsne_v1'
SOURCE = ROOT / 'PPT/胡丹-0917-QG-HRQ-Query正向案例与tSNE版.pptx'
OUTPUT = ROOT / 'PPT/胡丹-0917-QG-HRQ-TIGER对比版.pptx'
WORK = ROOT / 'qg_prqk/outputs/figures/qg_hrq_tiger_0917_v1'
NAMES = ('TIGER', 'QG-HRQ')
COLORS = ('#9BAFC4', '#3977B8')


def full_codes(sid: np.ndarray) -> np.ndarray:
    """Encode all three semantic positions, never compare S3 tokens alone."""
    if sid.ndim != 2 or sid.shape[1] != 3 or sid.dtype.kind not in 'iu':
        raise ValueError('SID 须为三列整数')
    if np.any(sid < 0) or np.any(sid >= 512):
        raise ValueError('SID 超出 512×3 码本范围')
    return (sid[:, 0].astype(np.int64) * 512 + sid[:, 1]) * 512 + sid[:, 2]


def collision_pairs(groups: np.ndarray, sid: np.ndarray) -> int:
    """Count pairs that share a geographic group and a complete semantic SID."""
    _, counts = np.unique(groups.astype(np.int64) * 512**3 + full_codes(sid), return_counts=True)
    return int(np.sum(counts * (counts - 1) // 2))


def checked(path: Path, sources: dict, expected: str | None = None) -> Path:
    actual = sha256_file(path)
    if expected and actual != expected:
        raise ValueError(f'来源哈希不一致：{path}')
    sources[str(path.relative_to(ROOT))] = actual
    return path


def frozen(directory: Path, name: str, sources: dict) -> Path:
    manifest = json.loads(checked(directory / 'manifest.json', sources).read_text())
    if str(manifest['status']).lower() != 'completed' or not (directory / '_SUCCESS').is_file():
        raise ValueError(f'来源阶段未完成：{directory}')
    return checked(directory / name, sources, manifest['artifacts'][name]['sha256'])


def query_statistics(sids: list[np.ndarray], sources: dict) -> dict:
    """Reuse the frozen effective edges and the original query-equal macro definition."""
    nodes = pq.read_table(frozen(GRAPH, 'query_nodes.parquet', sources),
                          columns=['node_row', 'query_id', 'normalized_query', 'query_count', 'supervision_depth'],
                          filters=[('supervision_depth', 'in', [1, 2])]).to_pylist()
    index = {n['node_row']: n for n in nodes}
    edges_path = frozen(GRAPH, 'query_poi_edges_s1_s2.parquet', sources)
    metrics, cases = {}, []
    rng = np.random.default_rng(42)
    for depth in (1, 2):
        edges = pq.read_table(edges_path, columns=['node_row', 'poi_row_index', 'pair_count', 'edge_probability'],
                              filters=[('supervision_depth', '=', depth), ('layer', '=', depth)])
        node, poi, count, probability = [edges[k].to_numpy() for k in edges.column_names]
        order = np.argsort(node, kind='stable')
        node, poi, count, probability = [a[order] for a in (node, poi, count, probability)]
        cuts = np.r_[0, np.flatnonzero(np.diff(node)) + 1, len(node)]
        codes = [prefix_codes(s, depth) for s in sids]
        buckets = [np.bincount(c, minlength=512**depth) for c in codes]
        queries = []
        for start, stop in zip(cuts[:-1], cuts[1:]):
            rows, mass = poi[start:stop], count[start:stop]
            if len(rows) < 2:
                continue
            q = index[int(node[start])]
            if len(np.unique(rows)) != len(rows) or mass.sum() > q['query_count']:
                raise ValueError('Query 边重复或计数不守恒')
            if not np.allclose(mass / mass.sum(), probability[start:stop], atol=1e-12):
                raise ValueError('Query 概率与原始有效边不一致')
            queries.append(dict(query=q['normalized_query'], query_id=q['query_id'],
                                depth=depth, query_count=q['query_count'], poi_rows=rows.tolist(),
                                pair_counts=mass.tolist(), retained_orders=int(mass.sum()),
                                before=distribution(codes[0][rows], mass, buckets[0]),
                                after=distribution(codes[1][rows], mass, buckets[1])))
        delta = np.asarray([q['after']['top_prefix_share'] - q['before']['top_prefix_share'] for q in queries])
        metrics[f'D{depth}'] = dict(
            multi_target_queries=len(queries),
            before={k: float(np.mean([q['before'][k] for q in queries])) for k in METRICS},
            after={k: float(np.mean([q['after'][k] for q in queries])) for k in METRICS},
            outcomes=dict(improved=int((delta > 1e-12).sum()), unchanged=int((abs(delta) <= 1e-12).sum()),
                          worsened=int((delta < -1e-12).sum())))
        cases.append(select_case(queries, rng, 'improved'))
        print(f'D{depth}: {json.dumps(metrics[f"D{depth}"], ensure_ascii=False)}', flush=True)
    # Keep the five queries already shown in the user's latest deck.
    old = json.loads(checked(OLD_QUERY / 'analysis.json', sources).read_text())
    projections = deepcopy(old['projection_cases'])
    for case in projections:
        for i, mode in enumerate(('before', 'after')):
            codes = prefix_codes(sids[i], 2)
            rows = [m['poi_row_index'] for m in case['members']]
            counts = [m['pair_count'] for m in case['members']]
            case[mode] = distribution(codes[rows], counts, np.bincount(codes, minlength=512**2))
            for member in case['members']:
                member[f'{mode}_prefix'] = sids[i][member['poi_row_index'], :2].tolist()
    rows = sorted({r for c in cases for r in c['poi_rows']})
    meta = pq.read_table(AFTER / 'poi_geo_metadata.parquet',
                         columns=['poi_row_index', 'poi_id', 'displayname', 'category'],
                         filters=[('poi_row_index', 'in', rows)]).to_pylist()
    meta = {m['poi_row_index']: m for m in meta}
    for c in cases:
        c['members'] = [{**meta[r], 'pair_count': n, 'before_prefix': sids[0][r, :c['depth']].tolist(),
                         'after_prefix': sids[1][r, :c['depth']].tolist()}
                        for r, n in zip(c.pop('poi_rows'), c.pop('pair_counts'))]
    return dict(metrics=metrics, cases=cases, projection_cases=projections,
                selection='Tree: seed42 random within improved D1/D2, 3..4 targets; t-SNE: unchanged five queries from uploaded deck. Examples are not population estimates.',
                contract='Train-only dominant-category effective edges; pair-count weighted within query; query-equal macro. No GID/C/Dedup.')


def analyze() -> tuple[dict, list[np.ndarray]]:
    sources = {}
    manifest = json.loads(checked(TIGER / 'sid_manifest.json', sources,
                                  'ba6738dd72d9eb6ab91470586809930d7a0cfa14a99a81b9c94e22add8caf52e').read_text())
    tiger = np.load(checked(TIGER / 'sid_codes.npy', sources,
                            'c6072dbeb622eae94f28ee28b0065693840680305c8d70a0d501172ec9d61c6d'))
    qg = np.load(frozen(AFTER, 'poi_sid_s1_s2_s3.npy', sources))
    selected = np.load(frozen(AFTER, 'selected_poi_rows.npy', sources))
    if tiger.shape != (716245, 3) or qg.shape != tiger.shape or not np.array_equal(selected, np.arange(len(qg))):
        raise ValueError('目录规模或 QG 行映射不一致')
    meta_path = frozen(AFTER, 'poi_geo_metadata.parquet', sources)
    metadata = pq.read_table(meta_path, columns=['poi_row_index', 'poi_id', 'category', 'category_code', 'gid6'])
    ids_path = checked((TIGER / manifest['poi_ids']['path']).resolve(), sources, manifest['poi_ids']['sha256'])
    qg_ids = metadata['poi_id'].to_pylist()
    with ids_path.open() as f:
        tiger_ids = [json.loads(line) for line in f]
    if tiger_ids != qg_ids or metadata['poi_row_index'].to_pylist() != list(range(len(qg))):
        raise ValueError('TIGER/QG POI ID 行序不一致')
    sids = [tiger, qg]
    fine = np.asarray(metadata['category'].to_pylist())
    _, gid = np.unique(np.asarray(metadata['gid6'].to_pylist()), return_inverse=True)
    purity, local = {}, {}
    hard = pq.read_table(frozen(AFTER, 'hard_entity_edges.parquet', sources),
                         columns=['poi_row_index', 'neighbor_poi_row_index', 'composite_similarity'])
    src, dst, weights = [hard[k].to_numpy() for k in hard.column_names]
    for name, sid in zip(NAMES, sids):
        levels = []
        for depth in (1, 2, 3):
            keys = prefix_codes(sid, depth) if depth < 3 else full_codes(sid)
            _, groups, sizes = np.unique(keys, return_inverse=True, return_counts=True)
            levels.append(_weighted_purity(groups, sizes, fine))
        purity[name] = levels
        codes = full_codes(sid)
        mask = codes[src] == codes[dst]
        local[name] = dict(collision_pairs=collision_pairs(gid, sid),
                           hard_weighted_collision_rate=float(weights[mask].sum() / weights.sum()))
    query = query_statistics(sids, sources)
    # Retain the original geographic cohort: no new selection on TIGER outcomes.
    old_geo = ROOT / 'qg_prqk/outputs/figures/qg_prqk_s3_local_v1/analysis.json'
    geo = json.loads(checked(old_geo, sources).read_text())['cases']
    for case in geo:
        rows = [m['poi_row_index'] for m in case['members']]
        for i, name in enumerate(NAMES):
            case[name] = dict(collision_pairs=collision_pairs(np.zeros(len(rows), dtype=int), sids[i][rows]))
            for m in case['members']:
                m[name] = sids[i][m['poi_row_index']].tolist()
    return dict(sources=sources, poi_count=len(qg), id_order_verified=True,
                purity=purity, local=local, geographic_cases=geo, query=query,
                limits=['Independent methods, not a Query/Geo-only ablation.',
                        'Local collisions compare full SID within common GID6; hard edges are the frozen QG training-edge cohort.',
                        'Geographic examples retain original POIs; equal S3 values across different S1/S2 are not collisions.']), sids


def plot_setup() -> None:
    os.environ.setdefault('MPLCONFIGDIR', str(ROOT / 'qg_prqk/outputs/cache/matplotlib'))
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import font_manager, pyplot as plt
    path = ROOT / 'qg_prqk/outputs/inputs/fonts/NotoSansCJKsc-Regular.otf'
    font_manager.fontManager.addfont(str(path))
    plt.rcParams.update({'font.family': font_manager.FontProperties(fname=str(path)).get_name(),
                         'font.size': 17, 'axes.unicode_minus': False})


def projection(data: dict, sids: list[np.ndarray], work: Path) -> None:
    """Use BGE bucket means in one common space for the two different quantizers."""
    from matplotlib import pyplot as plt
    from sklearn.decomposition import PCA
    from sklearn.manifold import TSNE
    from threadpoolctl import threadpool_limits
    cases = data['query']['projection_cases']
    rows = np.asarray([m['poi_row_index'] for c in cases for m in c['members']])
    labels = np.concatenate([np.full(len(c['members']), i) for i, c in enumerate(cases)])
    embeddings = np.load(EMBEDDING / 'embeddings.npy', mmap_mode='r')
    checked(EMBEDDING / 'manifest.json', data['sources'])
    parts, used_hash = [], hashlib.sha256()
    for sid in sids:
        codes = prefix_codes(sid, 2)
        wanted, inverse = np.unique(codes[rows], return_inverse=True)
        selected = np.flatnonzero(np.isin(codes, wanted))
        # Only the full buckets of the displayed prefixes are read, in catalog order.
        values = np.asarray(embeddings[selected], dtype=np.float32)
        used_hash.update(selected.tobytes())
        used_hash.update(values.tobytes())
        means = np.asarray([values[codes[selected] == c].mean(axis=0) for c in wanted])
        means /= np.maximum(np.linalg.norm(means, axis=1, keepdims=True), 1e-12)
        parts.append(means[inverse])
    unique, inverse = np.unique(np.concatenate(parts), axis=0, return_inverse=True)
    with threadpool_limits(limits=4):
        reduced = PCA(n_components=min(30, len(unique) - 1), random_state=42).fit_transform(unique)
        tsne = TSNE(n_components=2, perplexity=min(10, (len(unique) - 1)/3),
                    max_iter=1000, init='pca', learning_rate='auto', random_state=42, n_jobs=4)
        xy = tsne.fit_transform(reduced)[inverse].reshape(2, len(rows), 2)
    colors = ['#4878A8', '#E69F45', '#58A57D', '#B36B96', '#8D80BB']
    lo, hi = xy.min(axis=(0, 1)), xy.max(axis=(0, 1))
    margin = np.maximum((hi-lo)*.08, 1)
    for i, mode in enumerate(('tiger', 'qg')):
        fig, ax = plt.subplots(figsize=(5.8, 4.25), dpi=200)
        for label, color in enumerate(colors):
            points, counts = np.unique(xy[i, labels == label], axis=0, return_counts=True)
            ax.scatter(*points.T, s=45*counts, c=color, alpha=.85, edgecolors='white', linewidth=.6)
        ax.set_xlim(lo[0]-margin[0], hi[0]+margin[0])
        ax.set_ylim(lo[1]-margin[1], hi[1]+margin[1])
        ax.set_aspect('equal', adjustable='box')
        ax.axis('off')
        fig.subplots_adjust(left=0, right=1, top=1, bottom=0)
        fig.savefig(work / f'tsne_{mode}.png', facecolor='white')
        plt.close(fig)
    data['projection'] = dict(poi_rows=rows.tolist(), labels=labels.tolist(), coordinates=xy.tolist(), colors=colors,
                              seed=42, kl_divergence=float(tsne.kl_divergence_), selected_embedding_sha256=used_hash.hexdigest(),
                              representation='Full-catalog S1/S2 bucket mean of raw BGE-M3, unit normalized; common 1024D space.',
                              fit='Joint unique-vector PCA+tSNE, shared limits, identical vectors at identical coordinates; no jitter.',
                              reason='TIGER learned 256D codewords and QG 1024D codewords cannot be directly concatenated for a joint projection.')


def figures(data: dict, work: Path) -> None:
    from matplotlib import pyplot as plt
    fig, ax = plt.subplots(figsize=(12.2, 5.8))
    for i, name in enumerate(NAMES):
        values = [v['poi_weighted_top1_share']*100 for v in data['purity'][name]]
        bars = ax.bar(np.arange(3)+(i-.5)*.30, values, .30, label=name, color=COLORS[i])
        ax.bar_label(bars, labels=[f'{v:.2f}' for v in values], padding=6, fontsize=18)
    ax.set_xticks(np.arange(3), ['S1', 'S1 / S2', 'S1 / S2 / S3'])
    ax.set_ylim(0, 110)
    ax.set_yticks([0, 25, 50, 75, 100])
    ax.set_ylabel('Micro Purity (%)')
    ax.legend(frameon=False, loc='upper left', ncol=2)
    ax.spines[['top', 'right']].set_visible(False)
    fig.tight_layout()
    fig.savefig(work/'prefix_purity.png', dpi=180)
    plt.close(fig)
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.8))
    for ax, key, title, multiplier in zip(axes, ['collision_pairs','hard_weighted_collision_rate'],
                                         ['同一 GID6 内完整 SID 同码 POI 对', 'QG 固定困难边：完整 SID 同码率 (%)'], [1,100]):
        values = [data['local'][n][key]*multiplier for n in NAMES]
        bars = ax.bar(NAMES, values, width=.5, color=COLORS)
        ax.bar_label(bars, labels=[f'{v:,.0f}' if multiplier==1 else f'{v:.2f}' for v in values], padding=8, fontsize=20)
        ax.set_ylim(0, max(values)*1.24)
        ax.set_title(title, fontsize=18, pad=16)
        ax.spines[['top','right']].set_visible(False)
    fig.tight_layout(w_pad=3)
    fig.savefig(work/'local_overview.png', dpi=180)
    plt.close(fig)
    draw_local_cases(data, work)


def draw_local_cases(data: dict, work: Path) -> None:
    """Plot complete-SID equality at unchanged physical coordinates."""
    from matplotlib import pyplot as plt
    for j, case in enumerate(data['geographic_cases'], 1):
        members = case['members']
        xy = local_coordinates(np.asarray([m['lng'] for m in members]), np.asarray([m['lat'] for m in members]))
        center = (xy.min(axis=0)+xy.max(axis=0))/2
        span = max(float(np.ptp(xy[:,0])),float(np.ptp(xy[:,1]))*1.8,40)
        fig, axes = plt.subplots(1,2,figsize=(12,4.35),sharex=True,sharey=True)
        for ax, name in zip(axes, NAMES):
            codes = full_codes(np.asarray([m[name] for m in members]))
            _, labels = np.unique(codes,return_inverse=True)
            colors = plt.get_cmap('tab10')(labels)
            for left,right in same_code_pairs(codes):
                ax.plot(xy[[left,right],0],xy[[left,right],1],color=colors[left],alpha=.6,linewidth=1.3,zorder=1)
            for i, point in enumerate(xy):
                ax.scatter(*point,s=105,color=colors[i],edgecolor='white',linewidth=.8,zorder=3)
                dist = np.linalg.norm(xy-point,axis=1)
                dist[i] = np.inf
                delta = point-xy[int(np.argmin(dist))]
                offset = (8 if delta[0]>=0 else -17,10 if delta[1]>=0 else -20)
                if point[1]<center[1]-span*.25:
                    offset=(offset[0],10)
                ax.annotate(str(i+1),point,xytext=offset,textcoords='offset points',fontsize=17)
            ax.set_xlim(center[0]-span*.62,center[0]+span*.62)
            ax.set_ylim(center[1]-span*.34,center[1]+span*.34)
            ax.set_aspect('equal',adjustable='box')
            ax.set_title(name,fontsize=20,pad=12)
            ax.set_xlabel('东西向距离（米）',fontsize=16)
            ax.spines[['top','right']].set_visible(False)
        axes[0].set_ylabel('南北向距离（米）',fontsize=16)
        fig.subplots_adjust(left=.075,right=.975,top=.85,bottom=.18,wspace=.20)
        fig.savefig(work/f'local_case_{j}.png',dpi=180)
        plt.close(fig)


def replace_text(shape, text: str) -> None:
    """Preserve the user's first paragraph and run formatting."""
    frame = shape.text_frame
    template = deepcopy(frame.paragraphs[0]._p)
    for p in list(frame._txBody.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}p')):
        frame._txBody.remove(p)
    for line in text.split('\n'):
        element = deepcopy(template)
        frame._txBody.append(element)
        p = frame.paragraphs[-1]
        runs = list(p.runs)
        if runs:
            runs[0].text = line
            for run in runs[1:]:
                element.remove(run._r)
        else:
            p.add_run().text = line


def replace_picture(slide, index: int, path: Path) -> None:
    old = slide.shapes[index]
    new = slide.shapes.add_picture(str(path), old.left, old.top, width=old.width, height=old.height)
    old._element.addprevious(new._element)
    old._element.getparent().remove(old._element)


def query_tree(slide, cases: list[dict]) -> None:
    """Redraw only the existing tree body, using its editable typography and positions."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    template = deepcopy(slide.shapes[13]._element)
    for shape in list(slide.shapes)[3:]:
        slide.shapes._spTree.remove(shape._element)
    replace_text(slide.shapes[0], 'Query 关联：相关目标共享前缀的案例')
    replace_text(slide.shapes[1], 'TIGER')
    for header in (slide.shapes[1], slide.shapes[2]):
        header.top = Inches(1.40)

    def text(value, x, y, w, h, size=16, align=PP_ALIGN.LEFT):
        el=deepcopy(template)
        slide.shapes._spTree.insert_element_before(el,'p:extLst')
        shape=slide.shapes[-1]
        # Cloning cNvPr IDs must not introduce duplicates.
        shape._element.xpath('.//p:cNvPr')[0].set('id',str(max(s.shape_id for s in slide.shapes)+1))
        shape.left,shape.top,shape.width,shape.height=[Inches(v) for v in (x,y,w,h)]
        replace_text(shape,value)
        for p in shape.text_frame.paragraphs:
            p.alignment=align
            for run in p.runs:
                run.font.size=Pt(size)
        return shape

    def line(x1,y1,x2,y2,weight=1):
        shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*[Inches(v) for v in (x1,y1,x2,y2)])
        shape.line.color.rgb=RGBColor.from_string('64748B')
        shape.line.width=Pt(weight)

    text('D1 / D2 改善案例；组内固定 seed=42 抽取，不代表全量或单模块贡献。',.62,.88,12.06,.4,18)
    for row,case in enumerate(cases):
        top=2.0+row*2.55
        text(f"例  D{case['depth']}：{case['query']}（{'S1' if case['depth']==1 else 'S1/S2'}）",.69,top,7,.38,18)
        text(f"主前缀占比 {case['before']['top_prefix_share']:.1%} → {case['after']['top_prefix_share']:.1%}",8.12,top,4.55,.38,18,PP_ALIGN.RIGHT)
        for col,mode in enumerate(('before','after')):
            x=.7+col*6.18
            groups={}
            for m in case['members']:
                groups.setdefault(tuple(m[f'{mode}_prefix']),[]).append(m)
            groups=sorted(groups.items(),key=lambda item:(-sum(m['pair_count'] for m in item[1]),item[0]))
            cursor,positions=0,[]
            for code,members in groups:
                members=sorted(members,key=lambda m:m['poi_row_index'])
                ys=[top+.56+(cursor+j)*.53 for j in range(len(members))]
                positions.append((code,members,ys,sum(ys)/len(ys)))
                cursor+=len(members)
            root_y=top+.56+(cursor-1)*.53/2
            for code,members,ys,center in positions:
                share=sum(m['pair_count'] for m in members)/case['retained_orders']
                line(x+.17,root_y,x+1.05,center,1+share*2)
                for m,y in zip(members,ys):
                    line(x+2.18,center,x+2.65,y)
                text('/'.join(map(str,code))+f'\n{share:.1%}',x+1.02,center-.25,1.24,.55,16,PP_ALIGN.CENTER)
                for m,y in zip(members,ys):
                    # Keep the full name in analysis.json, but fit the editable tree label.
                    name = m['displayname']
                    if len(name) > 20:
                        name = name[:17] + '…'
                    text(f"{name}  ×{m['pair_count']}",x+2.7,y-.23,3,.55,16)
            dot=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+.11),Inches(root_y-.06),Inches(.12),Inches(.12))
            dot.fill.solid()
            dot.fill.fore_color.rgb=RGBColor.from_string('1A3A6B')
            dot.line.fill.background()


def update_deck(source: Path, output: Path, work: Path, data: dict) -> None:
    from pptx import Presentation
    source_hash=sha256_file(source)
    prs=Presentation(source)
    if len(prs.slides)!=26 or 'Query' not in prs.slides[11].shapes[0].text:
        raise ValueError('请使用用户更新的 26 页 0917 PPT')
    query_tree(prs.slides[11],data['query']['cases'])
    slide=prs.slides[12]
    table=slide.shapes[2].table
    for depth in (1,2):
        metric=data['query']['metrics'][f'D{depth}']
        for i,mode in enumerate(('before','after')):
            v=metric[mode]
            values=[f"D{depth} / {'S1' if depth==1 else 'S1/S2'}\n{NAMES[i]}",
                    f"{metric['multi_target_queries']:,}",f"{v['top_prefix_share']:.2%}",
                    f"{v['prefix_count']:.2f}",f"{v['entropy_bits']:.3f}",f"{v['expected_bucket_size']:,.2f}"]
            for cell,value in zip(table.rows[1+(depth-1)*2+i].cells,values):
                replace_text(cell,value)
    delta=[(data['query']['metrics'][f'D{i}']['after']['top_prefix_share']-data['query']['metrics'][f'D{i}']['before']['top_prefix_share'])*100 for i in (1,2)]
    replace_text(slide.shapes[3],f'主前缀目标占比：D1 {delta[0]:+.2f}pp；D2 {delta[1]:+.2f}pp（QG-HRQ − TIGER）')
    replace_text(slide.shapes[5],'两种方法的 Train 结构对比；需结合桶大小，不作为 Query 单项贡献或检索收益。')
    slide=prs.slides[13]
    replace_text(slide.shapes[0],'Query 关联结构：S1/S2 前缀语义 t-SNE')
    replace_text(slide.shapes[1],'原五组 Query、同一批 POI；用前缀桶的 BGE 均值做联合投影，seed=42。')
    replace_text(slide.shapes[2],'TIGER')
    replace_picture(slide,3,work/'tsne_tiger.png')
    replace_picture(slide,5,work/'tsne_qg.png')
    for i,c in enumerate(data['query']['projection_cases']):
        replace_text(slide.shapes[7+2*i],f"{c['query']}\n{c['before']['prefix_count']} → {c['after']['prefix_count']} 个前缀")
    replace_text(slide.shapes[16],'颜色表示 Query，点面积表示 POI 数；同前缀重合，无抖动。选例不代表全量。')
    replace_picture(prs.slides[14],1,work/'prefix_purity.png')
    slide=prs.slides[15]
    replace_text(slide.shapes[0],'局部实体区分：TIGER 与 QG-HRQ')
    replace_text(slide.shapes[1],'同一 716,245 POI 库、同一 GID6；比较完整 SID，非单独 S3 / Geo 消融。')
    replace_picture(slide,2,work/'local_overview.png')
    for i,case in enumerate(data['geographic_cases']):
        slide=prs.slides[16+i]
        replace_text(slide.shapes[0],slide.shapes[0].text.replace('S3 局部编码','局部实体编码'))
        replace_text(slide.shapes[1],f"同一批 POI、真实坐标；GID6：{case['gid6']}")
        replace_text(slide.shapes[2],'图内同色 / 连线 = 同完整 SID；编号 = 同一 POI')
        replace_picture(slide,3,work/f'local_case_{i+1}.png')
        for col,name in enumerate(NAMES):
            replace_text(slide.shapes[4+col],f"完整 SID 同码 POI 对数：{case[name]['collision_pairs']}")
    prs.save(output)
    validate_deck(source, output, work, source_hash, set(range(11, 18)))


def validate_deck(source: Path, output: Path, work: Path, source_hash: str,
                  changed: set[int]) -> None:
    """Check untouched pages and preserve all original notes and template parts."""
    from pptx import Presentation
    result=Presentation(output)
    original=Presentation(source)
    if len(result.slides)!=26 or sha256_file(source)!=source_hash:
        raise ValueError('源文件或页数发生变化')
    for i,(a,b) in enumerate(zip(original.slides,result.slides)):
        if i not in changed and etree.tostring(a._element,method='c14n')!=etree.tostring(b._element,method='c14n'):
            raise ValueError(f'非目标页 {i+1} 被修改')
        if a.has_notes_slide!=b.has_notes_slide or (a.has_notes_slide and
            etree.tostring(a.notes_slide._element,method='c14n')!=etree.tostring(b.notes_slide._element,method='c14n')):
            raise ValueError('备注不应改变')
    with ZipFile(source) as before,ZipFile(output) as after:
        if after.testzip() is not None:
            raise ValueError('PPT ZIP 损坏')
        for name in before.namelist():
            if name.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/')):
                a,b=[etree.fromstring(z.read(name)) for z in (before,after)]
                if name.endswith('.rels'):
                    a[:]=sorted(a,key=lambda e:e.get('Id'))
                    b[:]=sorted(b,key=lambda e:e.get('Id'))
                if etree.tostring(a,method='c14n')!=etree.tostring(b,method='c14n'):
                    raise ValueError('母版、布局或主题改变')
    write_json_atomic(work/'ppt_manifest.json',dict(source=str(source),source_sha256=source_hash,
        output=str(output),output_sha256=sha256_file(output),slides=26,changed_slides=[i+1 for i in sorted(changed)],
        unchanged_slides=26-len(changed),source_preserved=True,notes_masters_layouts_themes_preserved=True), overwrite=True)


def main() -> int:
    parser=argparse.ArgumentParser(description='以 active TIGER 重算并替换 0917 第 12–18 页可视化，保留上传稿其余页面。')
    parser.add_argument('--source',type=Path,default=SOURCE,help='用户最新的 26 页 PPT')
    parser.add_argument('--output',type=Path,default=OUTPUT,help='新 PPT，禁止覆盖')
    parser.add_argument('--work-dir',type=Path,default=WORK,help='qg_prqk 内图与溯源产物目录')
    parser.add_argument('--reuse-analysis',action='store_true',help='复用同目录已完成分析和图片，只重新排版')
    args=parser.parse_args()
    try:
        if args.output.exists() or args.output.resolve()==args.source.resolve():
            raise ValueError('输出 PPT 已存在或与源文件相同')
        if not args.work_dir.resolve().is_relative_to(ROOT/'qg_prqk/outputs'):
            raise ValueError('图表必须保存在 qg_prqk/outputs 下')
        if args.reuse_analysis:
            data=json.loads((args.work_dir/'analysis.json').read_text())
        else:
            if args.work_dir.exists() and any(args.work_dir.iterdir()):
                raise ValueError('工作目录已存在；如分析已完成，可使用 --reuse-analysis')
            args.work_dir.mkdir(parents=True, exist_ok=True)
            data,sids=analyze()
            plot_setup()
            projection(data,sids,args.work_dir)
            figures(data,args.work_dir)
            data['artifacts']={p.name:sha256_file(p) for p in args.work_dir.glob('*.png')}
            data['implementation_sha256']=sha256_file(Path(__file__))
            write_json_atomic(args.work_dir/'analysis.json',data)
        for name,value in data['artifacts'].items():
            if sha256_file(args.work_dir/name)!=value:
                raise ValueError(f'图表发生变化：{name}')
        update_deck(args.source,args.output,args.work_dir,data)
        print(json.dumps(dict(output=str(args.output),slides=26,purity=data['purity'],local=data['local']),ensure_ascii=False))
    except (ValueError,OSError,KeyError) as error:
        parser.exit(2,f'TIGER 可视化更新失败：{error}\n')
    return 0
