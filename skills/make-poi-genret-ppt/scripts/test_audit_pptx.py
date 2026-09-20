"""Synthetic regression tests for black titles and functional flow nodes."""

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt

from audit_pptx import EXPECTED_HEIGHT, EXPECTED_WIDTH, _inspect


class StyleAuditTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="ppt-skill-test-")
        self.path = Path(self.directory.name) / "synthetic.pptx"

    def tearDown(self):
        self.directory.cleanup()

    def make_deck(self, title_color="000000", outline=None):
        prs = Presentation()
        prs.slide_width, prs.slide_height = EXPECTED_WIDTH, EXPECTED_HEIGHT
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(10), Inches(0.7))
        run = title.text_frame.paragraphs[0].add_run()
        run.text, run.font.name, run.font.size = "测试标题", "Microsoft YaHei", Pt(28)
        if title_color:
            run.font.color.rgb = RGBColor.from_string(title_color)
        if outline:
            node = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                         Inches(1), Inches(2), Inches(3), Inches(1))
            node.line.color.rgb = RGBColor.from_string(outline)
        prs.save(self.path)

    def test_black_title_and_gray_node_pass(self):
        self.make_deck(outline="CBD5E1")
        report, issues = _inspect(self.path, False)
        self.assertEqual(issues, [])
        self.assertEqual(report["unexpected_explicit_text_colors"], [])

    def test_navy_major_title_fails(self):
        self.make_deck(title_color="1A3A6B")
        report, issues = _inspect(self.path, False)
        self.assertEqual(len(report["nonblack_titles"]), 1)
        self.assertTrue(issues)

    def test_inherited_title_color_requires_explicit_black(self):
        self.make_deck(title_color=None)
        self.assertTrue(_inspect(self.path, False)[1])

    def test_old_reference_or_notes_only_keeps_original_style(self):
        self.make_deck(title_color="1A3A6B")
        self.assertEqual(_inspect(self.path, False, enforce_black_titles=False)[1], [])

    def test_black_outline_still_fails(self):
        self.make_deck(outline="000000")
        report, issues = _inspect(self.path, False)
        self.assertEqual(len(report["dark_outlined_boxes"]), 1)
        self.assertTrue(issues)


if __name__ == "__main__":
    unittest.main()
