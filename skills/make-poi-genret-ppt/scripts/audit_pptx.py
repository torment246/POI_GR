#!/usr/bin/env python3
"""Audit a PPTX against the Hu Dan 0727/0810 reporting guardrails."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, PP_PLACEHOLDER


ALLOWED_FONTS = {"微软雅黑", "Microsoft YaHei"}
ALLOWED_TEXT_COLORS = {"000000", "111827", "1A3A6B", "64748B", "FFFFFF"}
SUSPICIOUS_GLYPHS = set("□■▪▫▣☐☑✅❌◆◇●○▲▼▶◀")
BOX_SHAPES = {
    MSO_AUTO_SHAPE_TYPE.RECTANGLE,
    MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
}
DARK_OUTLINE_COLORS = {"000000", "111827", "1A1A1A", "1F1F1F"}
EXPECTED_WIDTH = 12_192_000
EXPECTED_HEIGHT = 6_858_000


def _rgb(font: object) -> str | None:
    try:
        color = font.color
        if color.type is not None and color.rgb is not None:
            return str(color.rgb)
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def _canonical_slide_hashes(prs: Presentation) -> list[str]:
    hashes: list[str] = []
    for slide in prs.slides:
        payload = etree.tostring(
            slide._element,
            method="c14n",
            exclusive=False,
            with_comments=True,
        )
        hashes.append(hashlib.sha256(payload).hexdigest())
    return hashes


def _inspect(path: Path, require_notes: bool, *,
             enforce_black_titles: bool = True) -> tuple[dict[str, object], list[str]]:
    issues: list[str] = []
    with ZipFile(path) as archive:
        corrupt_member = archive.testzip()
        if corrupt_member is not None:
            issues.append(f"PPTX ZIP 损坏：{corrupt_member}")

    prs = Presentation(path)
    fonts: Counter[str] = Counter()
    colors: Counter[str] = Counter()
    suspicious: list[dict[str, object]] = []
    dark_outlined_boxes: list[dict[str, object]] = []
    empty_notes: list[int] = []
    shape_counts: list[int] = []
    nonblack_titles: list[dict[str, object]] = []

    if (prs.slide_width, prs.slide_height) != (EXPECTED_WIDTH, EXPECTED_HEIGHT):
        issues.append(
            "页面尺寸不是 0727/0810 的 16:9 基线："
            f"{prs.slide_width}x{prs.slide_height}"
        )

    for slide_number, slide in enumerate(prs.slides, start=1):
        shape_counts.append(len(slide.shapes))
        if require_notes:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if not note:
                empty_notes.append(slide_number)
        for shape in slide.shapes:
            try:
                line_color = shape.line.color
                line_rgb = (
                    str(line_color.rgb)
                    if line_color.type is not None and line_color.rgb is not None
                    else None
                )
                if (
                    shape.auto_shape_type in BOX_SHAPES
                    and line_rgb in DARK_OUTLINE_COLORS
                ):
                    dark_outlined_boxes.append(
                        {
                            "slide": slide_number,
                            "name": shape.name,
                            "line_color": line_rgb,
                            "geometry": [
                                shape.left,
                                shape.top,
                                shape.width,
                                shape.height,
                            ],
                        }
                    )
            except (AttributeError, TypeError, ValueError):
                pass
            if not getattr(shape, "has_text_frame", False):
                continue
            runs = [r for p in shape.text_frame.paragraphs for r in p.runs if r.text.strip()]
            is_major_title = (
                shape.is_placeholder
                and shape.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
            ) or any(r.font.size is not None and r.font.size.pt >= 28 for r in runs)
            if enforce_black_titles and is_major_title and any(_rgb(r.font) != "000000" for r in runs):
                nonblack_titles.append({"slide": slide_number, "text": shape.text,
                                        "colors": [_rgb(r.font) for r in runs]})
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        fonts[run.font.name] += len(run.text)
                    color = _rgb(run.font)
                    if color:
                        colors[color] += len(run.text)
                    for character in run.text:
                        codepoint = ord(character)
                        if character in SUSPICIOUS_GLYPHS or codepoint >= 0x1F000:
                            suspicious.append(
                                {
                                    "slide": slide_number,
                                    "character": character,
                                    "codepoint": f"U+{codepoint:04X}",
                                    "context": run.text[:80],
                                }
                            )

    unexpected_fonts = sorted(set(fonts) - ALLOWED_FONTS)
    unexpected_colors = sorted(set(colors) - ALLOWED_TEXT_COLORS)
    if suspicious:
        issues.append(f"发现 {len(suspicious)} 个可能显示为方框的装饰字符")
    if dark_outlined_boxes:
        issues.append(f"发现 {len(dark_outlined_boxes)} 个深色描边矩形")
    if empty_notes:
        issues.append(f"以下页面备注为空：{empty_notes}")
    if nonblack_titles:
        issues.append(f"发现 {len(nonblack_titles)} 个未显式设为黑色的大标题")

    report: dict[str, object] = {
        "path": str(path.resolve()),
        "slides": len(prs.slides),
        "slide_size": [prs.slide_width, prs.slide_height],
        "shape_counts": shape_counts,
        "canonical_slide_hashes": _canonical_slide_hashes(prs),
        "fonts": dict(fonts),
        "unexpected_explicit_fonts": unexpected_fonts,
        "explicit_text_colors": dict(colors),
        "unexpected_explicit_text_colors": unexpected_colors,
        "suspicious_glyphs": suspicious,
        "dark_outlined_boxes": dark_outlined_boxes,
        "empty_notes": empty_notes,
        "nonblack_titles": nonblack_titles,
    }
    return report, issues


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="检查 PPTX 页面尺寸、字体、颜色、方框字符、备注和正文稳定性。"
    )
    parser.add_argument("pptx", type=Path, help="待检查的 PPTX")
    parser.add_argument("--reference", type=Path, help="可选的 0727/0810 参考 PPTX")
    parser.add_argument("--baseline", type=Path, help="仅备注修改前的 PPTX 基线")
    parser.add_argument(
        "--require-notes",
        action="store_true",
        help="要求每页均有非空演讲者备注",
    )
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    if not args.pptx.is_file():
        print(f"文件不存在：{args.pptx}", file=sys.stderr)
        return 2

    # Notes-only edits must not restyle an older deck to satisfy newer rules.
    report, issues = _inspect(args.pptx, args.require_notes, enforce_black_titles=not args.baseline)

    if args.reference:
        reference, reference_issues = _inspect(args.reference, False, enforce_black_titles=False)
        if report["slide_size"] != reference["slide_size"]:
            issues.append("页面尺寸与参考 PPTX 不一致")
        report["reference"] = {
            "path": reference["path"],
            "slide_size": reference["slide_size"],
            "audit_warnings": reference_issues,
        }

    if args.baseline:
        baseline, baseline_issues = _inspect(args.baseline, False, enforce_black_titles=False)
        if baseline_issues:
            issues.append(f"基线 PPTX 本身存在问题：{baseline_issues}")
        if report["slides"] != baseline["slides"]:
            issues.append("仅备注修改后页面数量发生变化")
        if report["shape_counts"] != baseline["shape_counts"]:
            issues.append("仅备注修改后页面对象数量发生变化")
        if report["canonical_slide_hashes"] != baseline["canonical_slide_hashes"]:
            issues.append("仅备注修改后幻灯片正文 XML 发生变化")
        report["baseline"] = baseline["path"]

    report["issues"] = issues
    report["status"] = "failed" if issues else "passed"
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if issues else 0


if __name__ == "__main__":
    raise SystemExit(main())
